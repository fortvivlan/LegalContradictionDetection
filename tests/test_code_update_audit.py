"""Small local fixtures for the legal-code coverage audit."""

import hashlib
import json

from docx import Document

from LCD.experiments.lawbook.series_03_code_update.audit_codes import (
    audit_codes,
    extract_mentions,
    federal_law_numbers,
    inventory_codes,
    representative_mentions,
)


def test_mentions_include_inflections_abbreviations_and_unnumbered_references():
    records = extract_mentions(
        "Согласно Федеральному закону от 07.02.2011 № 3-ФЗ полиция вправе действовать. "
        "Применяется закон №52–ФЗ. Иное предложение.\n"
        "На основании федерального закона назначен штраф.\n"
        "Указанный ФЗ действует."
    )
    assert len(records) == 4
    assert [record["federal_law_numbers"] for record in records] == [["3"], ["52"], [], []]
    assert not extract_mentions("ФЗУ и федеральный орган не являются ссылками.")
    assert federal_law_numbers("№ 115-ФЗ, N 115-ФЗ и № 109 ФЗ") == ["109", "115"]
    assert federal_law_numbers("п. 4.2 ст. 13 ФЗ «О правовом положении»") == []
    assert federal_law_numbers("ст. 13 ФЗ № 115") == []


def test_inventory_preserves_filename_edition_and_content_hash(tmp_path):
    path = tmp_path / "Федеральный закон N 109 ФЗ_редакция 01.10.2026.rtf"
    path.write_bytes(b"fixture")
    record = inventory_codes(tmp_path)[0]
    assert record["edition_label"] == "01.10.2026"
    assert record["federal_law_numbers"] == ["109"]
    assert record["sha256"] == hashlib.sha256(b"fixture").hexdigest()


def test_audit_full_documents_tables_repeats_and_no_appeals_or_mirrors(tmp_path):
    codes = tmp_path / "codes"
    codes.mkdir()
    (codes / "Закон N 115 ФЗ_редакция 01.10.2026.rtf").write_bytes(b"fixture")
    decisions = tmp_path / "decisions"
    decisions.mkdir()
    text = "Применяется Федеральный закон № 3-ФЗ."
    document = Document()
    document.add_paragraph(text)
    document.add_paragraph("ПОСТАНОВИЛ:")
    document.add_paragraph(text)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Учитывается 115-ФЗ."
    document.save(decisions / "decision.docx")
    (decisions / "source.txt").write_text(text, encoding="utf-8")
    appeals = decisions / "appeals"
    appeals.mkdir()
    document.save(appeals / "appeal.docx")
    output = tmp_path / "output"
    summary = audit_codes(codes, decisions, output)
    assert summary["documents_scanned"] == 1
    assert summary["sentences_with_mentions"] == 3
    assert summary["mentioned_federal_law_sentence_counts"] == {"3": 2, "115": 1}
    assert summary["candidate_absent_federal_law_numbers"] == ["3"]
    report = (output / "federal_law_sentences.txt").read_text(encoding="utf-8")
    assert report.count(text) == 1
    assert (output / "federal_law_sentences_full.txt").read_text(encoding="utf-8").count(text) == 2
    assert summary["representative_sentences"] == 2
    assert "decision.docx" in report
    assert "appeal.docx" not in report
    assert json.loads((output / "code_versions.json").read_text(encoding="utf-8"))["codes"]


def test_representatives_distinguish_dates_and_merge_numbered_and_named_aliases():
    sentences = [
        'Согласно Федеральному закону № 3-ФЗ действует правило.',
        'Согласно Федеральному закону от 8 января 1998 г. № 3-ФЗ "О наркотических средствах" действует правило.',
        'Федеральный закон от 08.01.1998 № 3-ФЗ действует.',
        'Федеральный закон от 07.02.2011 № 3-ФЗ «О полиции» действует.',
        'Федеральный закон «О полиции» действует.',
        'Федеральный закон «О беженцах» действует.',
        'В соответствии с Федеральным законом «О беженцах» действует правило.',
        'Указанный Федеральный закон действует.',
    ]
    records = [
        {"source": "decision.docx", "sentence_index": index, "sentence": sentence,
         "federal_law_numbers": federal_law_numbers(sentence)}
        for index, sentence in enumerate(sentences)
    ]
    representatives = representative_mentions(records)
    assert len(representatives) == 3
    assert representatives[0]["law_labels"] == ["3-ФЗ от 08.01.1998"]
    assert representatives[1]["law_labels"] == ["3-ФЗ от 07.02.2011"]
    assert representatives[2]["sentence"] == sentences[5]


def test_representatives_merge_date_typos_with_same_descriptive_title():
    sentences = [
        'Федеральный закон от 25.07.2002 № 115-ФЗ «О правовом положении иностранных граждан» действует.',
        'Федеральный закон от 25.07.2022 № 115-ФЗ «О правовом положении иностранных граждан» действует.',
        'Нарушение от 17.08.2023 касается закона № 115-ФЗ от 25.07.2002 г. «О правовом положении иностранных граждан».',
        'Федеральный закон «О правовом положении иностранных граждан» от 25.07.2014 № 115-ФЗ действует.',
        'Федеральный закон № 115-ФЗ и Указ главы от 01.01.2025 № 1 «Об установлении запрета» действуют.',
        'Применяется ФЗ-115 «О правилах пребывания иностранцев».',
    ]
    records = [
        {"source": "decision.docx", "sentence_index": index, "sentence": sentence,
         "federal_law_numbers": federal_law_numbers(sentence)}
        for index, sentence in enumerate(sentences)
    ]
    representatives = representative_mentions(records)
    assert len(representatives) == 1
    assert representatives[0]["sentence"] == sentences[0]
