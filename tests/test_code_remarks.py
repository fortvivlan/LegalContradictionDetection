"""Small public fixtures for remark boundaries, scope, links and exports."""

import csv
import hashlib
import json

from docx import Document
from openpyxl import load_workbook
import pytest

from LCD.experiments.lawbook.series_04_remarks.annotate_remarks import (
    annotate_records, entity, expand_numbers, export_records, extract_remarks,
    read_paragraphs, suggest_annotation,
)


def paragraphs(*texts):
    return [{"index": i, "text": text, "style": "Normal"}
            for i, text in enumerate(texts, 1)]


def test_numbered_notes_editorial_interruptions_and_internal_lists():
    fixture = paragraphs("Глава 12. Заголовок", "Статья 12.1. Заголовок",
                         "Примечания:", "Изменено примечание 1", "1. Первое правило:",
                         "1) условие;", "2) другое условие.", "Продолжение правила.",
                         "2. Второе правило.", "Статья 12.2. Другая статья", "Основная норма.")
    fixture[3]["style"] = "Информация о версии"
    records = extract_remarks(fixture)
    assert len(records) == 2
    assert records[0]["remark_text"] == "1. Первое правило:\n1) условие;\n2) другое условие.\nПродолжение правила."
    assert records[0]["paragraph_start"] == 5
    assert records[0]["paragraph_end"] == 8
    assert records[1]["remark_number"] == "2"
    assert "Изменено" not in records[0]["remark_text"]
    assert "Основная норма" not in records[1]["remark_text"]


def test_inline_plural_repealed_and_continuation():
    records = extract_remarks(paragraphs(
        "Статья 1.1. Заголовок", "Примечание. Утратило силу.",
        "Примечания: 1. Новое правило.", "Второй абзац.", "2. Другое правило."))
    assert len(records) == 3
    assert records[0]["remark_number"] == ""
    assert records[1]["remark_text"] == "1. Новое правило.\nВторой абзац."
    assert [r["block_index"] for r in records] == [1, 2, 2]
    assert suggest_annotation(records[0])["targets"] == []
    assert suggest_annotation(records[0])["category"] == 9


def test_docx_styles_superscript_and_table_order(tmp_path):
    document = Document()
    document.add_paragraph("Статья 6.13\u00a01. Заголовок")
    document.add_paragraph("Примечание. Правило:")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Продолжение в таблице."
    document.add_paragraph("Статья 6.14. Следующая статья")
    source = tmp_path / "fixture.docx"
    document.save(source)
    records = extract_remarks(read_paragraphs(source))
    assert records[0]["article"] == "6.13.1"
    assert records[0]["remark_text"] == "Примечание. Правило:\nПродолжение в таблице."


def test_exception_citations_are_not_application_targets():
    records = extract_remarks(paragraphs("Глава 1. Заголовок", "Статья 1.5. Заголовок",
        "Примечание. Положение части 3 настоящей статьи не распространяется на правонарушения, предусмотренные частями 3.1 - 3.4 статьи 8.2, главой 12 настоящего Кодекса."))
    annotation = suggest_annotation(records[0])
    assert annotation["category"] == 2
    assert [t["entity_id"] for t in annotation["targets"]] == ["koap_rf:article:1.5:part:3"]


def test_range_expansion_and_ambiguous_ranges():
    assert expand_numbers("1.1, 8 - 10 и 12") == ["1.1", "8", "9", "10", "12"]
    assert expand_numbers("3.1 — 3.4") == ["3.1", "3.2", "3.3", "3.4"]
    with pytest.raises(ValueError, match="Ambiguous"):
        expand_numbers("3 - 3.4")


def test_overrides_require_matching_text_and_remark_targets():
    records = extract_remarks(paragraphs("Статья 16.2. Заголовок",
                                        "Примечания:", "3. Примечание 2 не применяется."))
    digest = hashlib.sha256(records[0]["remark_text"].encode()).hexdigest()
    override = {records[0]["remark_id"]: {"text_sha256": digest, "category": 6,
               "targets": [entity("remark", article="16.2", remark_number="2")]}}
    annotated = annotate_records(records, override)
    assert annotated[0]["targets"][0]["entity_id"] == "koap_rf:article:16.2:remark:2"
    assert annotated[0]["targets"][0]["part"] == ""
    override[records[0]["remark_id"]]["text_sha256"] = "wrong"
    with pytest.raises(ValueError, match="Stale"):
        annotate_records(records, override)


def test_full_text_roundtrip_csv_excel_and_normalized_links(tmp_path):
    document = Document()
    document.add_paragraph("Статья 1.5. Заголовок")
    document.add_paragraph("Примечание. Положения частей 1 и 2 настоящей статьи уточняются:")
    document.add_paragraph('Текст с запятой, кавычками "пример" и условиями.')
    document.add_paragraph("Статья 1.6. Заголовок")
    document.add_paragraph("Примечание. Утратило силу.")
    source = tmp_path / "fixture.docx"
    document.save(source)
    records = annotate_records(extract_remarks(read_paragraphs(source)))
    output = tmp_path / "out"
    manifest = export_records(records, source, output)
    with (output / "koap_remarks.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    workbook = load_workbook(output / "koap_remarks_review.xlsx")
    assert workbook["Примечания"].max_row == 3
    assert workbook["Примечания"]["A2"].value == rows[0]["remark_text"] == records[0]["remark_text"]
    assert len(json.loads(rows[0]["targets_json"])) == 2
    assert json.loads(rows[1]["targets_json"]) == []
    assert rows[1]["location_entity_id"] == "koap_rf:article:1.6"
    assert workbook["Связи"].max_row == 3
    assert manifest["remarks"] == 2
    assert manifest["target_links"] == 2
    assert manifest["repealed_placeholders"] == 1
    assert "9 — Утратившее силу" in (output / "remark_categories.txt").read_text(encoding="utf-8")

