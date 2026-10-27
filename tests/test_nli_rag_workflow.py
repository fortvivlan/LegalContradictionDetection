from __future__ import annotations

import csv
import json
from pathlib import Path
import shutil

import numpy as np
import pytest
from docx import Document
from openpyxl import Workbook

from LCD.experiments.data_creation.series_01_nli_rag.workflow import (
    DatasetEncoder,
    apply_variant,
    digest,
    initialize_bundle,
    prepare_bundle,
    read_json,
    retrieve_premises,
    select_context,
    validate_bundle,
)
from LCD.experiments.data_creation.series_01_nli_rag.preparation import extract_hypotheses
from LCD.experiments.data_creation.series_01_nli_rag.review import export_split


def _write_docx(path: Path, operative: str, *, prefix: str = "Обстоятельства дела.") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(prefix)
    document.add_paragraph("ПОСТАНОВИЛ:")
    for paragraph in operative.split("\n"):
        document.add_paragraph(paragraph)
    document.save(path)


def _write_expert_table(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Сигнал / ошибка",
            "Тип сигнала / ошибки",
            "Добавляем в текущую выборку (октябрь 2026)",
            "Процент документов от общего количества синтетических документов текущей основной и отдельной выборок",
            "Комментарий",
            "Инструкция для внесения синтетической ошибки",
            "Промпт для поиска сигналов / ошибок",
            "Пример 1",
            "Пример 2",
        ]
    )
    sheet.append(["Неверная сумма", "Противоречие", "Да", "100%", "", "Измени сумму", "Проверь", "", ""])
    workbook.save(path)


def _corpus(tmp_path: Path) -> tuple[Path, Path]:
    root = tmp_path / "decisions"
    common = (
        "Признать лицо виновным по ч. 1 ст. 18.8 КоАП РФ и назначить штраф 2000 рублей.\n"
        "Штраф должен быть уплачен в течение 60 дней.\n"
        "Постановление может быть обжаловано в течение 10 суток.\n"
        "Судья И.И. Иванов"
    )
    definitions = [
        ("18.8/a.docx", common, "court-a", "case-1"),
        ("18.8/b.docx", common, "court-b", "case-2"),
        ("18.10/c.docx", common.replace("18.8", "18.10"), "court-c", "case-3"),
        ("18.15/d.docx", common.replace("18.8", "18.15"), "court-d", "case-4"),
    ]
    fields = ["filename", "group", "article", "source_url", "document_id", "court", "case_number"]
    rows = []
    for number, (relative, text, court, case) in enumerate(definitions, 1):
        _write_docx(root / "main" / relative, text)
        rows.append(
            {
                "filename": f"main/{relative}",
                "group": "main",
                "article": relative.split("/")[0],
                "source_url": f"https://example.test/{number}",
                "document_id": f"document-{number}",
                "court": court,
                "case_number": case,
            }
        )
    with (root / "provenance.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    table = tmp_path / "rules.xlsx"
    _write_expert_table(table)
    return root, table


class FakeEncoder:
    configuration = {
        "model": "fake-qwen",
        "revision": "0" * 40,
        "device": "cpu",
        "batch_size": 4,
        "normalize_embeddings": True,
        "dimension": 3,
    }

    @staticmethod
    def _vectors(texts):
        values = []
        for text in texts:
            code = sum(text.encode("utf-8"))
            values.append([1.0 + code % 5, 1.0 + code % 7, 1.0 + code % 11])
        return np.asarray(values, dtype="float32")

    def encode_documents(self, texts):
        return self._vectors(texts)

    def encode_queries(self, texts, instruction):
        assert instruction
        return self._vectors([instruction + text for text in texts])


def test_qwen_adapter_bounds_padded_tokens_and_restores_input_order() -> None:
    from types import SimpleNamespace

    calls = []

    class Model:
        def encode(self, texts, **kwargs):
            calls.append(texts)
            return np.asarray([[len(t), ord(t[0])] for t in texts], dtype="float32")

    native = SimpleNamespace(model=Model(), max_tokens=1000, batch_size=4,
                             configuration={"model": "fake"},
                             token_lengths=lambda texts: [len(t) for t in texts])
    encoder = DatasetEncoder(native, max_batch_tokens=16)
    texts = ["abc", "b" * 8, "cde", "d" * 8]
    actual = encoder.encode_documents(texts)
    expected = np.asarray([[len(t), ord(t[0])] for t in texts], dtype="float32")
    expected /= np.linalg.norm(expected, axis=1, keepdims=True)
    assert np.allclose(actual, expected)
    assert len(calls) == 2
    assert all(max(len(t) for t in batch) * len(batch) <= 16 for batch in calls)
    native.max_tokens = 2
    with pytest.raises(ValueError, match="no truncation"):
        encoder.encode_documents(["oversized"])


def test_atomic_checkpoint_retries_windows_sharing_lock(tmp_path, monkeypatch) -> None:
    from LCD.experiments.data_creation.series_01_nli_rag import _io

    source, target = tmp_path / "temp", tmp_path / "checkpoint.json"
    source.write_text("new", encoding="utf-8")
    target.write_text("old", encoding="utf-8")
    original_replace = Path.replace
    attempts = []

    def temporarily_locked(path, destination):
        attempts.append(path)
        if len(attempts) < 3:
            raise PermissionError("Windows reader holds a transient sharing lock")
        return original_replace(path, destination)

    monkeypatch.setattr(Path, "replace", temporarily_locked)
    monkeypatch.setattr(_io.time, "sleep", lambda seconds: None)
    _io.replace_checkpoint(source, target)
    assert len(attempts) == 3
    assert target.read_text(encoding="utf-8") == "new"


def _unit(number: int, article: str = "18.8") -> dict:
    return {
        "unit_id": f"unit-{number}",
        "text": f"Статья {article}, положение {number}.",
        "article": article,
        "part": "1",
        "point": None,
        "source": f"КоАП РФ: Статья {article} ч. 1",
        "rank": number,
        "score": 1.0 / number,
    }


class FakeStore:
    def __init__(self):
        self.manifest = {
            "source_sha256": "a" * 64,
            "index": {"query_instruction": "retrieve law", "revision": "0" * 40},
        }
        self.calls = []

    def exact_lookup(self, article, *, part=None, point=None):
        self.calls.append(("exact", article, part, point))
        return [_unit(1, article)]

    def search(self, hypothesis, *, top_k):
        self.calls.append(("search", hypothesis, top_k))
        return [_unit(number) for number in range(1, top_k + 1)]

    def search_many(self, hypotheses, *, top_k):
        self.calls.append(("search_many", tuple(hypotheses), top_k))
        return [[_unit(number) for number in range(1, top_k + 1)] for _ in hypotheses]


def _target_hypothesis(bundle: Path, source: dict) -> tuple[dict, str, str]:
    split_dir = bundle / source["split"]
    split_manifest = read_json(split_dir / "manifest.json")
    original = next(
        variant for variant in split_manifest["variants"]
        if variant["source_id"] == source["source_id"] and variant["variant_id"] == "original"
    )
    text = (split_dir / original["decision_text_path"]).read_text(encoding="utf-8")
    hypothesis = next(item for item in extract_hypotheses(text)[0] if "2000" in item["text"])
    hypothesis_id = "hyp_" + digest(
        [source["source_id"], hypothesis["start"], hypothesis["text"]]
    )[:20]
    return hypothesis, hypothesis_id, text


def _proposal(source: dict, hypothesis: dict, hypothesis_id: str, text: str, replacement: str) -> dict:
    start = text.index("2000", hypothesis["start"], hypothesis["end"])
    return {
        "source_id": source["source_id"],
        "rule_id": "xlsx_row_2",
        "status": "proposed",
        "target_hypothesis_id": hypothesis_id,
        "rationale": "Создаёт целевое расхождение суммы.",
        "edits": [{"start": start, "end": start + 4, "before": "2000", "after": replacement}],
    }


def test_initialize_is_deterministic_groups_duplicates_and_relocates(tmp_path: Path) -> None:
    root, table = _corpus(tmp_path)
    first, second = tmp_path / "bundle-one", tmp_path / "bundle-two"

    assert initialize_bundle(root, first, table, sizes=(2, 1, 1), seed=9)["sizes"] == {
        "train": 2,
        "val": 1,
        "test": 1,
    }
    initialize_bundle(root, second, table, sizes=(2, 1, 1), seed=9)
    first_manifest, second_manifest = read_json(first / "manifest.json"), read_json(second / "manifest.json")
    assert first_manifest["documents"] == second_manifest["documents"]
    duplicates = [
        document for document in first_manifest["documents"]
        if document["source_filename"].endswith(("a.docx", "b.docx"))
    ]
    assert {document["split"] for document in duplicates} == {"train"}

    relocated = tmp_path / "moved" / "portable-bundle"
    relocated.parent.mkdir()
    shutil.move(str(first), relocated)
    report = validate_bundle(relocated)
    assert {split: report[split]["documents"] for split in ("train", "val", "test")} == {
        "train": 2,
        "val": 1,
        "test": 1,
    }


def test_filter_refresh_excludes_old_signature_without_reassigning_ids(tmp_path: Path) -> None:
    from LCD.experiments.data_creation.series_01_nli_rag.curation import refresh_filter
    from LCD.experiments.data_creation.series_01_nli_rag.review import register_rows

    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=7)
    prepare_bundle(bundle, FakeStore(), FakeEncoder(), split="train", top_k=2, hard_limit=3)
    split_dir = bundle / "train"
    registry = read_json(split_dir / "registry.json")
    first = next(iter(registry["rows"].values()))
    ids_before = {key: row["row_id"] for key, row in registry["rows"].items()}
    text = (split_dir / first["decision_text_path"]).read_text(encoding="utf-8")
    signature = "Судья И.И. Иванов"
    start = text.index(signature)
    stale = {**first, "stable_key": "stale-bank-or-signature-row", "hypothesis": signature,
             "metadata": {**first["metadata"], "hypothesis_span": [start, start + len(signature)]}}
    stale.pop("row_id")
    allocated = register_rows(split_dir, [stale])[0]

    result = refresh_filter(bundle, split="train")

    after = read_json(split_dir / "registry.json")
    assert allocated in result["train"]["row_ids"]
    assert after["rows"]["stale-bank-or-signature-row"]["status"] == "excluded"
    assert all(after["rows"][key]["row_id"] == value for key, value in ids_before.items())
    assert register_rows(split_dir, [stale]) == [allocated]
    assert read_json(split_dir / "registry.json")["rows"]["stale-bank-or-signature-row"]["status"] == "excluded"
    validate_bundle(bundle)


def test_preparation_review_recovers_readable_hypotheses_from_legacy_errors(tmp_path: Path) -> None:
    from LCD.experiments.data_creation.series_01_nli_rag.curation import enrich_preparation_errors

    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=7)
    store = FakeStore()
    store.exact_lookup = lambda *args, **kwargs: [_unit(i) for i in range(1, 6)]
    prepare_bundle(bundle, store, FakeEncoder(), split="train", top_k=2, hard_limit=3)
    errors = read_json(bundle / "preparation_errors.json")
    assert errors
    for item in errors:
        for field in ("hypothesis", "source_filename", "decision_text_path", "citations"):
            item.pop(field, None)
    (bundle / "preparation_errors.json").write_text(json.dumps(errors), encoding="utf-8")

    enriched = enrich_preparation_errors(bundle)

    assert all(item["hypothesis"].startswith("Признать") for item in enriched)
    assert all(item["citations"] and item["source_filename"] for item in enriched)
    assert all((bundle / item["split"] / item["decision_text_path"]).is_file() for item in enriched)
    assert len({item["review_id"] for item in enriched}) == len(enriched)
    assert (bundle / "preparation_review.jsonl").is_file()


def test_retrieval_exact_first_cap_and_context_tie_behavior() -> None:
    store = FakeStore()
    results, unresolved = retrieve_premises(
        "Правило установлено ч. 1 ст. 18.8 КоАП РФ.", store, top_k=3, hard_limit=4
    )
    assert not unresolved
    assert [call[0] for call in store.calls] == ["exact", "search"]
    assert len(results) == 3
    assert results[0]["methods"] == ["exact", "faiss"]
    store.exact_lookup = lambda *args, **kwargs: [_unit(number) for number in range(1, 6)]
    with pytest.raises(ValueError, match="exceeding 4"):
        retrieve_premises(
            "Правило установлено ч. 1 ст. 18.8 КоАП РФ.",
            store,
            top_k=3,
            hard_limit=4,
            semantic=[],
        )

    sentences = [
        {"text": "Контекст позже.", "start": 30, "end": 45, "index": 2},
        {"text": "Целевая гипотеза.", "start": 10, "end": 28, "index": 1},
        {"text": "Контекст раньше.", "start": 0, "end": 9, "index": 0},
        {"text": "  КОНТЕКСТ РАНЬШЕ. ", "start": 50, "end": 69, "index": 3},
    ]
    context, audit = select_context(
        "целевая гипотеза.", sentences,
        np.asarray([[1, 0], [1, 0], [1, 0], [1, 0]], dtype="float32"),
        np.asarray([1, 0], dtype="float32"), count=5,
    )
    assert context == "Контекст раньше.\nКонтекст позже."
    assert [item["start"] for item in audit] == [0, 30]
    assert {item["rank"] for item in audit} == {1, 2}


@pytest.mark.parametrize("split", ["train", "val", "test"])
def test_equal_provision_text_at_distinct_citations_is_preserved(tmp_path, split) -> None:
    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=5)
    store = FakeStore()
    units = [{**_unit(i), "text": "Одинаковый текст разных частей нормы.", "part": str(i)}
             for i in (1, 2)]
    store.exact_lookup = lambda *args, **kwargs: []
    store.search_many = lambda hypotheses, **kwargs: [units for _ in hypotheses]
    prepare_bundle(bundle, store, FakeEncoder(), split=split, top_k=2, hard_limit=3)
    rows = list(read_json(bundle / split / "registry.json")["rows"].values())
    first_hypothesis = rows[0]["hypothesis_id"]
    selected = [row for row in rows if row["hypothesis_id"] == first_hypothesis]
    assert {r["unit_id"] for r in selected} == {"unit-1", "unit-2"}
    assert len({r["row_id"] for r in selected}) == 2
    validate_bundle(bundle)


def test_prepare_is_resumable_and_preserves_ids(tmp_path: Path) -> None:
    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=4)
    store, encoder = FakeStore(), FakeEncoder()

    first = prepare_bundle(bundle, store, encoder, split="val", top_k=2, hard_limit=3)
    registry_before = read_json(bundle / "val" / "registry.json")
    ids_before = {row["row_id"] for row in registry_before["rows"].values()}
    second = prepare_bundle(bundle, store, encoder, split="val", top_k=2, hard_limit=3)
    registry_after = read_json(bundle / "val" / "registry.json")

    assert first["prepared_variants"] == 1
    assert second["prepared_variants"] == 0
    assert ids_before and ids_before == {row["row_id"] for row in registry_after["rows"].values()}
    assert all(row_id.startswith("val_") for row_id in ids_before)
    assert len(read_json(bundle / "val" / "manifest.json")["prepared"]) == 1


def test_apply_variant_rejects_stale_edits_and_enforces_split_limits(tmp_path: Path) -> None:
    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=2)
    root_manifest = read_json(bundle / "manifest.json")
    train_source = next(document for document in root_manifest["documents"] if document["split"] == "train")
    hypothesis, hypothesis_id, text = _target_hypothesis(bundle, train_source)

    stale = _proposal(train_source, hypothesis, hypothesis_id, text, "1000")
    stale["edits"][0]["before"] = "9999"
    with pytest.raises(ValueError, match="stale"):
        apply_variant(bundle, stale)
    apply_variant(bundle, _proposal(train_source, hypothesis, hypothesis_id, text, "1000"))
    apply_variant(bundle, _proposal(train_source, hypothesis, hypothesis_id, text, "3000"))
    with pytest.raises(ValueError, match="Two distinct variants"):
        apply_variant(bundle, _proposal(train_source, hypothesis, hypothesis_id, text, "4000"))

    val_source = next(document for document in root_manifest["documents"] if document["split"] == "val")
    val_hypothesis, val_id, val_text = _target_hypothesis(bundle, val_source)
    apply_variant(bundle, _proposal(val_source, val_hypothesis, val_id, val_text, "1000"))
    with pytest.raises(ValueError, match="one active variant"):
        apply_variant(bundle, _proposal(val_source, val_hypothesis, val_id, val_text, "3000"))
    val_manifest = read_json(bundle / "val" / "manifest.json")
    original = next(variant for variant in val_manifest["variants"] if variant["variant_id"] == "original")
    assert original["active"] is False


@pytest.mark.parametrize("exclude_first", [False, True])
def test_train_preparation_merges_exact_pairs_from_identical_variant(tmp_path: Path, exclude_first) -> None:
    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=3)
    store, encoder = FakeStore(), FakeEncoder()
    prepare_bundle(bundle, store, encoder, split="train", top_k=2, hard_limit=3)
    split_dir = bundle / "train"
    registry_before = read_json(split_dir / "registry.json")
    first_key = next(iter(registry_before["rows"]))
    if exclude_first:
        registry_before["rows"][first_key]["status"] = "excluded"
        (split_dir / "registry.json").write_text(json.dumps(registry_before), encoding="utf-8")
    count_before = len(registry_before["rows"])
    manifest = read_json(split_dir / "manifest.json")
    original = manifest["variants"][0]
    clone_id = "syn_identical"
    clone_relative = f"decisions/variants/{clone_id}.txt"
    (split_dir / "decisions" / "variants").mkdir()
    shutil.copyfile(split_dir / original["decision_text_path"], split_dir / clone_relative)
    manifest["variants"].append(
        {**original, "variant_id": clone_id, "decision_text_path": clone_relative,
         "rule_id": "xlsx_row_2", "target_hypothesis_id": "test-only"}
    )
    (split_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

    assert prepare_bundle(bundle, store, encoder, split="train", top_k=2, hard_limit=3)[
        "prepared_variants"
    ] == 1
    registry_after = read_json(split_dir / "registry.json")
    assert len(registry_after["rows"]) == count_before
    if exclude_first:
        assert registry_after["rows"][first_key]["status"] == "excluded"
    clone_checkpoint = read_json(split_dir / "manifest.json")["prepared"][
        original["source_id"] + ":" + clone_id
    ]
    assert clone_checkpoint["row_ids"]
    assert set(clone_checkpoint["row_ids"]) <= {
        row["row_id"] for row in registry_after["rows"].values()
    }
    merged = [
        row for row in registry_after["rows"].values()
        if clone_id in row["metadata"].get("variant_ids", [])
    ]
    assert merged
    assert all(row["decision_text_path"].startswith("decisions/originals/") for row in merged)
    assert all("xlsx_row_2" in row["metadata"]["scenario_rule_ids"] for row in merged)


def test_apply_variant_rejects_target_units_not_retrieved_for_hypothesis(tmp_path: Path) -> None:
    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=5)
    prepare_bundle(bundle, FakeStore(), FakeEncoder(), split="train", top_k=2, hard_limit=3)
    source = next(
        document for document in read_json(bundle / "manifest.json")["documents"]
        if document["split"] == "train"
    )
    hypothesis, hypothesis_id, text = _target_hypothesis(bundle, source)
    proposal = _proposal(source, hypothesis, hypothesis_id, text, "1000")
    proposal["target_unit_ids"] = ["invented-unit"]

    with pytest.raises(ValueError, match="not retrieved candidates"):
        apply_variant(bundle, proposal)


class EmptyStore(FakeStore):
    def exact_lookup(self, article, *, part=None, point=None):
        return []

    def search_many(self, hypotheses, *, top_k):
        return [[] for _ in hypotheses]


def test_failed_preparation_is_checkpointed_and_errors_are_preserved(tmp_path: Path) -> None:
    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=6)

    first = prepare_bundle(bundle, EmptyStore(), FakeEncoder(), split="val", top_k=2, hard_limit=3)
    errors_before = read_json(bundle / "preparation_errors.json")
    second = prepare_bundle(bundle, EmptyStore(), FakeEncoder(), split="val", top_k=2, hard_limit=3)

    assert first["prepared_variants"] == second["prepared_variants"] == 0
    assert errors_before == read_json(bundle / "preparation_errors.json")
    manifest = read_json(bundle / "val" / "manifest.json")
    assert not manifest["prepared"]
    assert len(manifest["prepared_failed"]) == 1
    assert next(iter(manifest["prepared_failed"].values()))["errors"][0]["reason"] == "no_pairs"


def test_validate_checks_registry_semantics_and_csv_identity(tmp_path: Path) -> None:
    root, table = _corpus(tmp_path)
    bundle = tmp_path / "bundle"
    initialize_bundle(root, bundle, table, sizes=(2, 1, 1), seed=7)
    prepare_bundle(bundle, FakeStore(), FakeEncoder(), split="val", top_k=2, hard_limit=3)
    split_dir = bundle / "val"
    registry = read_json(split_dir / "registry.json")
    row = next(iter(registry["rows"].values()))
    row.update(status="proposed", label="contradiction", relevance="unrelated",
               error_type="Неверная сумма")
    (split_dir / "registry.json").write_text(json.dumps(registry), encoding="utf-8")
    with pytest.raises(ValueError, match="must be related"):
        validate_bundle(bundle)

    row["relevance"] = "related"
    (split_dir / "registry.json").write_text(json.dumps(registry), encoding="utf-8")
    export_split(split_dir)
    csv_path = split_dir / "val.csv"
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    rows[0]["hypothesis"] += " изменено"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=rows[0])
        writer.writeheader()
        writer.writerows(rows)
    with pytest.raises(ValueError, match="canonical registry"):
        validate_bundle(bundle)
