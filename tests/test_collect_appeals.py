from __future__ import annotations

import csv
from pathlib import Path

import pytest
from docx import Document

from LCD.experiments.data_creation.collect_appeals import (
    PublishedDecision,
    Verification,
    article_pattern,
    parse_sudact_page,
    persist_decision,
    search_sudact,
    verify_appeal,
)


def _decision(body: str) -> PublishedDecision:
    return PublishedDecision(
        document_id="ABC123",
        url="https://sudact.ru/regular/doc/ABC123/",
        title="Решение № 12-1/2025",
        court="Районный суд",
        case_number="12-1/2025",
        paragraphs=(body,),
        source_html="<html>source</html>",
    )


def _accepted_body(operative: str) -> str:
    return (
        "Судья, рассмотрев жалобу Иванова на постановление начальника отдела "
        "по делу об административном правонарушении, предусмотренном ч. 1 ст. 18.8 КоАП РФ. "
        "УСТАНОВИЛ: Постановлением начальника отдела Иванов признан виновным в совершении "
        "административного правонарушения по ч. 1 ст. 18.8 КоАП РФ и ему назначено "
        "административное наказание. Суд исследовал доказательства и доводы сторон. "
        "Руководствуясь статьями 30.6-30.8 КоАП РФ, суд РЕШИЛ: " + operative + " " + "Обоснование. " * 30
    )


def test_parser_limits_text_to_published_body() -> None:
    page = (
        '<h1>Решение № 12-1/2025 от 1 января 2025 г.</h1>'
        '<div class="b-justice"><a>Районный суд</a> - Административные правонарушения</div>'
        '<hr class="hr-h1">Текст решения<br>Вторая строка'
        '<!--AdFox START-->реклама и связанный документ'
    )
    parsed = parse_sudact_page(page, "ABC123", "https://example.test")
    assert parsed is not None
    assert parsed.paragraphs == ("Текст решения", "Вторая строка")
    assert "реклама" not in parsed.text
    assert parsed.court == "Районный суд - Административные правонарушения"


def test_verification_accepts_final_chapter_30_appeal() -> None:
    decision = _decision(_accepted_body(
        "Постановление начальника отдела по ч. 1 ст. 18.8 КоАП РФ оставить без изменения, "
        "жалобу Иванова оставить без удовлетворения."
    ))
    result = verify_appeal(decision, "18.8")
    assert result.accepted
    assert result.disposition == "upheld_denied"
    assert all((result.chapter30, result.challenged_ruling, result.complaint,
                result.article, result.disposition_evidence))


def test_verification_accepts_singular_chapter_30_citation() -> None:
    body = _accepted_body(
        "Постановление по ст. 18.8 КоАП РФ отменить, производство по делу прекратить."
    ).replace("статьями 30.6-30.8 КоАП РФ", "ст. 30.7 КоАП РФ")
    result = verify_appeal(_decision(body), "18.8")
    assert result.accepted
    assert result.disposition == "revoked_and_closed"


def test_verification_rejects_original_and_interim_acts() -> None:
    original = _decision(
        "Суд рассмотрел дело об административном правонарушении по ст. 18.8 КоАП РФ. "
        "ПОСТАНОВИЛ: признать Иванова виновным и назначить штраф. " + "Текст. " * 100
    )
    interim = _decision(
        "Суд рассмотрел жалобу на постановление по ст. 18.8 КоАП РФ. "
        "Руководствуясь ст. 30.3 КоАП РФ, ОПРЕДЕЛИЛ: возвратить жалобу заявителю. "
        + "Текст. " * 100
    )
    assert not verify_appeal(original, "18.8").accepted
    assert not verify_appeal(interim, "18.8").accepted


def test_verification_rejects_article_mentioned_only_as_history() -> None:
    body = (
        "Судья, рассмотрев жалобу на постановление инспектора по делу об административном "
        "правонарушении. УСТАНОВИЛ: Постановлением инспектора Иванов признан виновным "
        "по ст. 18.8 КоАП РФ и ему назначено административное наказание. Ранее Иванов "
        "привлекался по ст. 18.20 КоАП РФ. Руководствуясь ст. 30.6-30.8 КоАП РФ, "
        "суд РЕШИЛ: Постановление по ст. 18.8 КоАП РФ оставить без изменения, "
        "жалобу без удовлетворения. " + "Текст. " * 100
    )
    assert not verify_appeal(_decision(body), "18.20").accepted


def test_verification_rejects_bare_chapter_30_article() -> None:
    body = _accepted_body(
        "Постановление по ст. 18.8 КоАП РФ оставить без изменения, "
        "жалобу без удовлетворения."
    ).replace("статьями 30.6-30.8 КоАП РФ", "статьями 30.6-30.8")
    assert not verify_appeal(_decision(body), "18.8").accepted


def test_verification_rejects_complaint_against_prior_review_decision() -> None:
    body = (
        "Судья, рассмотрев жалобу Иванова на решение районного суда, принятое по жалобе "
        "на вынесенное в отношении Иванова постановление по делу об административном "
        "правонарушении, предусмотренном ст. 18.8 КоАП РФ. УСТАНОВИЛ: Постановлением "
        "инспектора Иванов признан виновным по ст. 18.8 КоАП РФ и ему назначено "
        "административное наказание. Руководствуясь ст.ст. 30.6-30.8 КоАП РФ, суд "
        "РЕШИЛ: Постановление отменить, производство по делу прекратить. "
        + "Текст. " * 100
    )
    result = verify_appeal(_decision(body), "18.8")
    assert not result.accepted
    assert "prior review decision" in result.reason


def test_verification_rejects_prior_review_named_before_original_in_opening() -> None:
    body = (
        "Судья, рассмотрев жалобу защитника на решение районного суда от 18 декабря "
        "2025 года по делу, предусмотренному ч. 2 ст. 18.11 КоАП РФ. "
        "УСТАНОВИЛ: постановлением начальника отдела Иванов признан виновным "
        "по ч. 2 ст. 18.11 КоАП РФ. Защитник ранее обратился с жалобой на "
        "постановление начальника отдела. Районный суд оставил постановление "
        "без изменения. Руководствуясь ст. 30.7 КоАП РФ, суд РЕШИЛ: "
        "решение районного суда и постановление начальника отдела по "
        "ч. 2 ст. 18.11 КоАП РФ оставить без изменения, жалобу без удовлетворения. "
        + "Обоснование. " * 30
    )
    result = verify_appeal(_decision(body), "18.11")
    assert not result.accepted
    assert "prior review decision" in result.reason


def test_article_pattern_does_not_confuse_adjacent_articles() -> None:
    pattern = article_pattern("18.8")
    assert pattern.search("часть 1 статьи 18.8 КоАП РФ")
    assert not pattern.search("ст. 18.80 КоАП РФ")
    assert not pattern.search("ст. 118.8 КоАП РФ")


@pytest.mark.parametrize("article, text", [
    ("18.11", "постановление от 18.11.2025 по ст. 18.8 КоАП РФ"),
    ("18.12", "постановление от 18.12. 2025 г. по ст. 20.21 КоАП РФ"),
    ("18.19", "постановление №3/18.19-33/2025 по ч.1 ст. 8.8 КоАП РФ"),
    ("18.11", "ст. 18.11.1 КоАП РФ"),
    ("18.11", "ст. 18.110 КоАП РФ"),
])
def test_article_pattern_rejects_dates_identifiers_and_subarticles(article, text) -> None:
    assert not article_pattern(article).search(text)


@pytest.mark.parametrize("article, marker", [
    ("18.11", "18.11.2025"),
    ("18.12", "18.12. 2025 г."),
    ("18.19", "№3/18.19-33/2025"),
])
def test_verification_rejects_non_citation_in_current_complaint(article, marker) -> None:
    body = _accepted_body(
        "Постановление по ч. 1 ст. 18.8 КоАП РФ оставить без изменения, "
        "жалобу без удовлетворения."
    ).replace("на постановление начальника отдела", f"на постановление начальника отдела {marker}")
    assert not verify_appeal(_decision(body), article).accepted


def test_verification_rejects_opening_typo_conflicting_with_operative_charge() -> None:
    body = _accepted_body(
        "Постановление по ч. 1 ст. 18.8 КоАП РФ оставить без изменения, "
        "жалобу без удовлетворения."
    ).replace("ч. 1 ст. 18.8 КоАП РФ", "ч. 1 ст. 18.18 КоАП РФ", 1)
    result = verify_appeal(_decision(body), "18.18")
    assert not result.accepted
    assert "conflicts" in result.reason


def test_verification_flags_matching_labels_with_conflicting_reasoned_offence() -> None:
    body = _accepted_body(
        "Постановление по ч. 3 ст. 18.19 КоАП РФ оставить без изменения, "
        "жалобу без удовлетворения."
    ).replace("ч. 1 ст. 18.8 КоАП РФ", "ч. 3 ст. 18.19 КоАП РФ")
    body = body.replace(
        "Суд исследовал доказательства и доводы сторон.",
        "Факт совершения Ивановым административного правонарушения, "
        "предусмотренного частью 3 статьей 18.9 КоАП РФ, подтвержден."
    )
    result = verify_appeal(_decision(body), "18.19")
    assert not result.accepted
    assert "substantive offence" in result.reason


def test_search_polls_and_deduplicates() -> None:
    responses = iter([
        '{"status": "new", "search_status": "new"}',
        '{"content": "<a href=\\"/regular/doc/ABC123/\\">one</a>'
        '<a href=\\"/regular/doc/ABC123/\\">same</a>'
        '<a href=\\"/regular/doc/XYZ789/\\">two</a>"}',
    ])
    assert search_sudact(lambda _: next(responses), "18.8 жалобу", 1, poll_delay=0) == [
        "ABC123", "XYZ789",
    ]


def test_persist_writes_source_text_docx_and_provenance_values(tmp_path: Path) -> None:
    decision = _decision(_accepted_body(
        "Постановление отменить, производство по делу прекратить."
    ))
    verification = Verification(
        True, "accepted", "revoked_and_closed", "chapter", "ruling",
        "complaint", "article", "disposition",
    )
    row = persist_decision(tmp_path, "main", "18.8", decision, verification)
    for field in ("filename", "source_html", "source_text"):
        assert (tmp_path / row[field]).is_file()
    assert Document(tmp_path / row["filename"]).paragraphs
    assert row["group"] == "main"
    assert row["article"] == "18.8"
    assert len(row["content_sha256"]) == 64
