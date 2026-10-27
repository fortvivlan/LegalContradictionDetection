"""Small DOCX fixtures and mock embeddings; no private data or model downloads."""

import csv
import json

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
import pytest

from LCD.shared.coderag.corpus import (
    Paragraph, extract_docx, extract_paragraphs, export_corpus, read_paragraphs,
)
from LCD.shared.coderag.index import (
    QUERY_INSTRUCTION, CodeRagStore, QwenEncoder, build_index, query_text,
)


def fixture_paragraphs(*texts):
    return [Paragraph(i, text) for i, text in enumerate(texts, 1)]


def fixture_docx(tmp_path):
    document = Document()
    document.add_paragraph("Раздел I. Общие положения")
    document.add_paragraph("Глава 1. Правила")
    document.add_paragraph("Статья 1.3. Полномочия")
    document.add_paragraph("1. К ведению относится:")
    document.add_paragraph("1) первый вопрос;")
    document.add_paragraph("2) второй вопрос.")
    document.add_paragraph("2. Самостоятельное правило.")
    document.add_paragraph("Примечание.  Исходный\u00a0текст.")
    source = tmp_path / "КоАП_редакция 01.10.2026.docx"
    document.save(source)
    return source


def test_mixed_branches_article_points_and_nested_lists():
    result = extract_paragraphs(fixture_paragraphs(
        "Раздел IV. Производство", "Глава 26. Доказательства",
        "Статья 26.4. Экспертиза", "1. Первое правило.",
        "2. В определении указываются:", "1) основания;", "2) вопросы:",
        "а) первый;", "б) второй.", "Кроме того, требуется подпись.",
        "3. Следующее правило.", "Статья 26.5. Другое", "Нужно выяснить:",
        "1) событие;", "2) участников.", "Статья 26.6. Без частей", "Правило -",
        "влечет наложение штрафа."))
    units = {r["source"]: r for r in result["units"]}
    assert len(units) == 7
    assert units["п. 1 ч. 2 ст. 26.4"]["text"] == (
        "В определении указываются:\nоснования;\nКроме того, требуется подпись.")
    assert "а) первый;\nб) второй." in units["п. 2 ч. 2 ст. 26.4"]["text"]
    assert units["п. 1 ст. 26.5"]["text"] == "Нужно выяснить:\nсобытие;"
    assert units["ст. 26.6"]["text"] == "Правило -\nвлечет наложение штрафа."
    assert units["ч. 1 ст. 26.4"]["section"] == "4"


def test_point_specific_continuation_and_ambiguous_tail():
    base = ["Статья 4.3. Правило", "1. Обстоятельства:", "1) первое;", "6) опьянение."]
    result = extract_paragraphs(fixture_paragraphs(
        *base, "Судья может не признать данное обстоятельство отягчающим."))
    assert "Судья" not in result["units"][0]["text"]
    assert "Судья" in result["units"][1]["text"]
    with pytest.raises(ValueError, match="Ambiguous post-point"):
        extract_paragraphs(fixture_paragraphs(*base, "Неясное общее правило."))


def test_amendments_following_heading_and_point_inheritance():
    paragraphs = fixture_paragraphs(
        "Раздел I. Общие", "Глава 1. Правила",
        "Информация об изменениях:", "Статья 1.3 изменена законом.",
        "Статья 1.3. Полномочия", "ГАРАНТ:", "См. комментарии к статье 1.3",
        "Информация об изменениях:", "Часть 1 изменена законом.",
        "1. Относится:", "1) первый вопрос;", "Информация об изменениях:",
        "Часть 1 дополнена пунктом 2.", "2) второй вопрос.", "2. Иное.")
    for i in (3, 8, 12):
        p = paragraphs[i]
        paragraphs[i] = Paragraph(p.index, p.text, "Информация о версии")
    for i in (5, 6):
        p = paragraphs[i]
        paragraphs[i] = Paragraph(p.index, p.text, "Комментарий")
    units = extract_paragraphs(paragraphs)["units"]
    assert "Статья 1.3 изменена" in units[0]["changes"]
    assert "Часть 1 изменена" in units[0]["changes"]
    assert "дополнена пунктом 2" not in units[0]["changes"]
    assert "дополнена пунктом 2" in units[1]["changes"]
    assert "Часть 1 изменена" not in units[2]["changes"]
    assert all("ГАРАНТ" not in r["text"] and "См." not in r["text"] for r in units)


def test_notes_verbatim_numbered_continuations_and_inactive():
    result = extract_paragraphs(fixture_paragraphs(
        "Раздел II. Особенная часть", "Глава 12. Правила", "Статья 12.1. Норма",
        "1. Основное правило.", "2. Утратила силу.",
        "Примечания:", "1.  Текст\u00a0с пробелами.", "  Продолжение.  ", "1) список;",
        "2. Утратило силу.", "3. Иное правило.",
        "Статья 12.2. Утратила силу с 2020 года.",
        "Статья 12.3. Не применяется с 2018 года.",
        "Статья 12.4. Иное", "Правило о законе, который утратил силу.",
        "Примечание.  Исходный\u00a0текст.", "Президент", "Российской Федерации"))
    assert len(result["units"]) == 2
    assert len(result["notes"]) == 3
    assert result["notes"][0]["text"] == "1.  Текст\u00a0с пробелами.\n  Продолжение.  \n1) список;"
    assert result["notes"][0]["position"] == "ст. 12.1 глава 12 раздел 2"
    assert result["notes"][-1]["text"] == "Примечание.  Исходный\u00a0текст."
    assert all("Президент" not in r["text"] for r in result["units"] + result["notes"])


def test_docx_raw_text_styles_superscript_and_table_order(tmp_path):
    document = Document()
    document.styles.add_style("Комментарий", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Раздел I. Общие")
    document.add_paragraph("Глава 6. Нормы")
    heading = document.add_paragraph("Статья 6.13")
    heading.add_run("\u00a01").font.superscript = True
    heading.add_run(". Правило")
    document.add_paragraph("ГАРАНТ: комментарий", style="Комментарий")
    document.add_paragraph("1.3-1. Действие -")
    document.add_paragraph("влечет штраф.")
    document.add_paragraph("Примечание.  Текст\u00a0с пробелами.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Продолжение\tв таблице."
    source = tmp_path / "КоАП.docx"
    document.save(source)
    result = extract_docx(source)
    assert result["units"][0]["source"] == "ч. 1.3-1 ст. 6.13.1"
    assert result["units"][0]["text"] == "Действие -\nвлечет штраф."
    assert result["notes"][0]["text"] == "Примечание.  Текст\u00a0с пробелами.\nПродолжение\tв таблице."
    assert read_paragraphs(source)[3].style == "Комментарий"


def test_csv_roundtrip_audit_and_deterministic_ids(tmp_path):
    source = fixture_docx(tmp_path)
    output = tmp_path / "output"
    manifest = export_corpus(source, output)
    original = extract_docx(source)
    with (output / "koap_units.csv").open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        assert reader.fieldnames == ["text", "source", "hierarchy", "changes"]
        rows = list(reader)
    assert rows[0]["text"] == "К ведению относится:\nпервый вопрос;"
    assert manifest["edition"] == "01.10.2026"
    assert manifest["counts"]["units"] == 3
    records = [json.loads(line) for line in (output / "units.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["unit_id"] for r in records] == [r["unit_id"] for r in original["units"]]
    audit = json.loads((output / "extraction_audit.json").read_text(encoding="utf-8"))
    assert all(p["unit_ids"] for p in audit["paragraphs"] if p["kind"] in {"legal_text", "note"})


class FakeEncoder:
    max_tokens = 1000
    instruction = QUERY_INSTRUCTION
    configuration = {"model": "fake", "revision": "fixture", "query_instruction": QUERY_INSTRUCTION,
                     "precision": "float32", "normalize_embeddings": True}

    def token_lengths(self, texts):
        return [len(text) for text in texts]

    def encode_documents(self, texts):
        self.documents = texts
        return [[1, 0], [0, 1], [-1, 0]]

    def encode_query(self, text):
        return [1, 0]


def test_index_roundtrip_row_identity_lookup_and_cosine(tmp_path):
    pytest.importorskip("faiss")
    output = tmp_path / "output"
    export_corpus(fixture_docx(tmp_path), output)
    encoder = FakeEncoder()
    manifest = build_index(output, encoder=encoder)
    assert manifest["index"]["vectors"] == 3
    assert encoder.documents == ["К ведению относится:\nпервый вопрос;",
                                 "К ведению относится:\nвторой вопрос.", "Самостоятельное правило."]
    store = CodeRagStore(output, encoder=encoder)
    hits = store.search("Первый вопрос", top_k=20)
    assert len(hits) == 3 and hits[0]["source"] == "п. 1 ч. 1 ст. 1.3"
    assert [r["score"] for r in hits] == pytest.approx([1, 0, -1])
    assert len(store.exact_lookup("1.3", part="1")) == 2
    assert len(store.exact_lookup("1.3", point="2")) == 1
    assert store.exact_lookup("99") == []
    wrong_encoder = FakeEncoder()
    wrong_encoder.configuration = {**wrong_encoder.configuration, "query_instruction": "Changed instruction"}
    with pytest.raises(ValueError, match="query_instruction"):
        CodeRagStore(output, encoder=wrong_encoder)
    with pytest.raises(ValueError, match="fresh"):
        export_corpus(fixture_docx(tmp_path), output)
    with (output / "faiss_index/units.jsonl").open("a", encoding="utf-8") as stream:
        stream.write("{}\n")
    with pytest.raises(ValueError, match="hash mismatch"):
        CodeRagStore(output, encoder=encoder)


def test_index_rejects_overflow_and_csv_tampering(tmp_path):
    pytest.importorskip("faiss")
    output = tmp_path / "output"
    export_corpus(fixture_docx(tmp_path), output)
    encoder = FakeEncoder()
    encoder.max_tokens = 1
    with pytest.raises(ValueError, match="no truncation.*п. 1 ч. 1 ст. 1.3"):
        build_index(output, encoder=encoder)
    assert not (output / "faiss_index").exists()
    with (output / "koap_units.csv").open("a", encoding="utf-8") as stream:
        stream.write("tampered")
    with pytest.raises(ValueError, match="hash mismatch"):
        build_index(output, encoder=encoder)


def test_exact_query_prompt_and_no_document_prompt():
    assert QUERY_INSTRUCTION == (
        "Given a statement from a Russian court decision, retrieve provisions of Russian law "
        "that regulate the same legal issue, including rules that may support or contradict "
        "the statement, while prioritizing matching actors, actions, conditions, and procedural stages.")
    assert query_text("Постановление суда.") == f"Instruct: {QUERY_INSTRUCTION}\nQuery: Постановление суда."
    encoder = object.__new__(QwenEncoder)
    encoder.instruction, encoder.max_tokens, encoder.batch_size = QUERY_INSTRUCTION, 1000, 8

    class Model:
        def tokenizer(self, texts, **kwargs):
            assert not kwargs["truncation"]
            return {"input_ids": [[0] * len(t) for t in texts]}

        def encode(self, texts, **kwargs):
            self.texts, self.kwargs = texts, kwargs
            return [[1, 0] for _ in texts]

    encoder.model = Model()
    encoder.encode_query("Судебное решение")
    assert encoder.model.texts == [query_text("Судебное решение")]
    assert encoder.model.kwargs["prompt"] == ""
    encoder.encode_documents(["Правовая норма"])
    assert encoder.model.texts == ["Правовая норма"]
    assert encoder.model.kwargs["prompt"] == ""
    encoder.max_tokens = 1
    with pytest.raises(ValueError, match="no truncation"):
        encoder.encode_query("Длинный запрос")
