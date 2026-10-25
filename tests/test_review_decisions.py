"""Small public-style fixtures for expert corpus curation; no private data."""

import csv
from pathlib import Path

from docx import Document
import pytest

from LCD.experiments.data_creation.review_decisions import (
    analyze_text, apply_review, create_review, entrepreneur_party_evidence,
)


GUILT = "Признать ФИО1 виновным в совершении правонарушения, предусмотренного ч. 1 ст. 18.8 КоАП РФ"
PENALTY = "назначить ему наказание в виде штрафа в размере 2000 рублей"
JOINED = GUILT + ", и " + PENALTY + " с административным выдворением за пределы РФ."
EMBEDDED = (
    "Согласно ч. 1 ст. 27.19.1 КоАП РФ, срок содержания иностранного гражданина "
    "в специальном учреждении, предусмотренном Федеральным законом от 25 июля "
    "2002 года N 115-ФЗ, не должен превышать девяносто суток."
)


def review(operative: str, opening: str = "") -> dict:
    return analyze_text(opening + "\nПОСТАНОВИЛ:\n" + operative)


def test_federal_law_context_and_independent_reference_in_same_document():
    assert "direct_federal_law" not in review(JOINED + "\n" + EMBEDDED)["reasons"]
    result = review(JOINED + "\n" + EMBEDDED + "\nВ силу п. 5 ст. 34 ФЗ содержать ФИО1 в центре.")
    assert "direct_federal_law" in result["reasons"]
    assert result["protected_federal_law_contexts"] == [EMBEDDED]


def test_payment_law_is_excluded_even_after_koap_payment_sentence():
    result = review(
        JOINED + '\nСогласно ст. 32.2 КоАП РФ штраф уплатить в срок. '
        'Сумма штрафа вносится через агента по Федеральному закону "О национальной платежной системе".'
    )
    assert "direct_federal_law" in result["reasons"]


def test_federal_legislation_and_references_before_operative_are_allowed():
    assert "direct_federal_law" not in review(
        JOINED + "\nИсполнить в порядке, предусмотренном федеральным законодательством.",
        "Применяется ст. 27 ФЗ от 15.08.1996 № 114-ФЗ.",
    )["reasons"]


@pytest.mark.parametrize("suffix", ["", "\nВыдворение не применять.", "\nБез административного выдворения."])
def test_separate_article_and_fine_without_expulsion_are_retained(suffix):
    assert "split_penalty" not in review(GUILT + ".\n" + PENALTY.capitalize() + "." + suffix)["reasons"]


@pytest.mark.parametrize("text", [
    GUILT + ".\n" + PENALTY.capitalize() + ".\nВыдворить ФИО1 за пределы РФ.",
    GUILT + ", и " + PENALTY + ".\nВыдворить ФИО1 за пределы РФ.",
    GUILT + ".\n" + PENALTY.capitalize() + " с административным выдворением за пределы РФ.",
])
def test_separated_imposed_penalties_are_excluded(text):
    assert "split_penalty" in review(text)["reasons"]


def test_joined_sentence_and_warning_of_future_expulsion():
    assert "split_penalty" not in review(JOINED)["reasons"]
    result = review(GUILT + ".\n" + PENALTY.capitalize() + ".\n"
                    "За уклонение предусмотрены штраф и принудительное выдворение.")
    assert "split_penalty" not in result["reasons"]


def test_separate_departure_explanation_and_cited_explanation():
    detached = "Иностранный гражданин обязан выехать из РФ в течение пяти дней."
    assert "split_departure_explanation" in review(JOINED + "\n" + detached)["reasons"]
    assert "split_departure_explanation" not in review(
        JOINED + "\nСогласно ч. 6 ст. 32.10 КоАП РФ " + detached.lower()
    )["reasons"]


def test_entrepreneur_party_versus_employer_or_statutory_reference():
    assert "entrepreneur_mention" in review(JOINED.replace("ФИО1", "ИП ФИО1", 1))["reasons"]
    assert "entrepreneur_mention" in review(JOINED, "Дело в отношении индивидуального предпринимателя ФИО1. УСТАНОВИЛ:")["reasons"]
    assert "entrepreneur_mention" not in review(JOINED, "ФИО1 работал у ИП ФИО2.")["reasons"]
    assert "entrepreneur_mention" not in review(JOINED, "По закону работодателями могут быть индивидуальные предприниматели.")["reasons"]
    employer_in_subject = JOINED.replace("ФИО1 виновным", "ФИО1, работающего у ИП ФИО2, виновным")
    assert "entrepreneur_mention" not in review(employer_in_subject)["reasons"]


def test_entrepreneur_after_guilt_and_citation():
    text = "Признать виновным по ч.1 ст.18.15 КоАП РФ индивидуального предпринимателя ФИО1 и назначить штраф."
    assert entrepreneur_party_evidence(text, text)


def test_multi_offences_not_penalty_mitigation_or_requalification():
    multiple = JOINED.replace("ч. 1 ст. 18.8", "ч. 1 ст. 18.8, ч. 1 ст. 18.10")
    assert "multiple_offences" in review(multiple)["reasons"]
    same_article = JOINED.replace("ч. 1 ст. 18.8", "ч. 1 ст. 18.8, ч. 2 ст. 18.8")
    assert "multiple_offences" in review(same_article)["reasons"]
    assert "multiple_offences" not in review(JOINED.replace("и назначить", "и с применением ст. 4.1 КоАП РФ назначить"))["reasons"]
    assert "multiple_offences" not in review("Переквалифицировать с ч. 4 ст. 18.8 КоАП РФ на ч. 1 ст. 18.8 КоАП РФ. " + JOINED)["reasons"]


def test_missing_section_is_pending_and_spacing_is_review_candidate():
    assert "missing_operative" in analyze_text("РЕШИЛ: отказать")["review_flags"]
    assert "possible_poor_spacing" in review(JOINED.replace("размере 2000", "размере2000"))["review_flags"]


def make_pool(tmp_path: Path) -> tuple[Path, Path]:
    root = tmp_path / "decisions"
    fields = ["filename", "source_text", "source_html", "source_url", "group", "article"]
    rows = []
    for group, stem, op in (("main", "keep", JOINED),
                            ("main", "exclude", JOINED + "\nВ силу ст. 27 ФЗ въезд не разрешается."),
                            ("domain_shift", "untouched", JOINED)):
        directory = root / group / "18.8"
        directory.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_paragraph("ПОСТАНОВИЛ:")
        doc.add_paragraph(op)
        doc.save(directory / f"{stem}.docx")
        for suffix in ("txt", "html"):
            (directory / f"{stem}.{suffix}").write_text(op, encoding="utf-8")
        rows.append(dict(filename=f"{group}/18.8/{stem}.docx", source_text=f"{group}/18.8/{stem}.txt",
                         source_html=f"{group}/18.8/{stem}.html", source_url="https://example.com/" + stem,
                         group=group, article="18.8"))
    with (root / "provenance.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return root, tmp_path / "review"


def test_apply_archives_bundle_preserves_provenance_and_other_pool(tmp_path):
    root, output = make_pool(tmp_path)
    before = (root / "provenance.csv").read_bytes()
    create_review(root, output)
    assert (root / "main/18.8/exclude.docx").is_file()
    result = apply_review(output / "plan.json")
    assert result["excluded"] == 1
    assert result["moved_files"] == 3
    assert (output / "before_review/provenance.csv").read_bytes() == before
    assert (output / "excluded/main/18.8/exclude.docx").is_file()
    assert not (root / "main/18.8/exclude.docx").exists()
    assert (root / "domain_shift/18.8/untouched.docx").is_file()
    with (root / "provenance.csv").open(encoding="utf-8") as stream:
        assert len(list(csv.DictReader(stream))) == 2
    with pytest.raises(FileExistsError):
        apply_review(output / "plan.json")


def test_source_hash_change_stops_apply_before_any_moves(tmp_path):
    root, output = make_pool(tmp_path)
    create_review(root, output)
    (root / "main/18.8/keep.txt").write_text("changed")
    with pytest.raises(ValueError, match="Source changed"):
        apply_review(output / "plan.json")
    assert (root / "main/18.8/exclude.docx").is_file()


def test_failed_move_rolls_back_documents_and_manifest(tmp_path, monkeypatch):
    import LCD.experiments.data_creation.review_decisions as module

    root, output = make_pool(tmp_path)
    create_review(root, output)
    original = (root / "provenance.csv").read_bytes()
    real_move = module.shutil.move
    calls = []

    def interrupted_move(source, destination):
        calls.append(source)
        if len(calls) == 2:
            raise OSError("injected move failure")
        return real_move(source, destination)

    monkeypatch.setattr(module.shutil, "move", interrupted_move)
    with pytest.raises(OSError, match="injected"):
        apply_review(output / "plan.json")
    assert (root / "main/18.8/exclude.docx").is_file()
    assert (root / "provenance.csv").read_bytes() == original


def test_unreviewed_ambiguous_cases_cannot_be_applied(tmp_path):
    root, output = make_pool(tmp_path)
    source = root / "main/18.8/keep.docx"
    doc = Document(source)
    doc.add_paragraph("Дела объединены в одно производство.")
    doc.save(source)
    create_review(root, output)
    with pytest.raises(ValueError, match="pending review"):
        apply_review(output / "plan.json")
