"""Rebuild the local validation CSV from balanced and expert-reviewed rows.

The expert workbook is authoritative for added pairs. Its row order and cell
text are preserved; importing this module never reads or writes local data.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


FIELDS = ("premise", "hypothesis", "source", "tag")
EXPERT_FIELDS = (*FIELDS, "document")
PROVENANCE_FIELDS = ("workbook_row", "document", *FIELDS)
LABELS = {"contradiction", "entailment", "not mentioned"}
DEFAULT_DIR = Path("local/data/classification")


def read_balanced_rows(path: Path) -> list[dict[str, str]]:
    """Read and validate a balanced four-column validation CSV."""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != list(FIELDS):
            raise ValueError(f"{path}: expected columns {FIELDS}")
        rows = list(reader)
    if not rows:
        raise ValueError(f"{path}: no validation rows")
    for row_number, row in enumerate(rows, 2):
        if any(not isinstance(row[field], str) or not row[field].strip()
               for field in FIELDS) or row["tag"] not in LABELS:
            raise ValueError(f"{path}: invalid row {row_number}")
    return rows


def read_expert_rows(path: Path) -> list[dict[str, str]]:
    """Read all checked pairs from the workbook's named review sheet."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Added validation pairs" not in workbook.sheetnames:
            raise ValueError(f"{path}: missing 'Added validation pairs' sheet")
        values = workbook["Added validation pairs"].values
        if next(values, None) != EXPERT_FIELDS:
            raise ValueError(f"{path}: expected columns {EXPERT_FIELDS}")
        rows = []
        for row_number, cells in enumerate(values, 2):
            if len(cells) != len(EXPERT_FIELDS):
                raise ValueError(f"{path}: invalid row {row_number}")
            row = dict(zip(EXPERT_FIELDS, cells))
            if any(not isinstance(row[field], str) or not row[field].strip()
                   for field in EXPERT_FIELDS):
                raise ValueError(f"{path}: missing text in row {row_number}")
            if row["tag"] != "not mentioned":
                raise ValueError(f"{path}: unexpected label in row {row_number}")
            rows.append({"workbook_row": str(row_number), **row})
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f"{path}: no expert-reviewed pairs")
    return rows


def build_validation(
    balanced_path: Path, expert_path: Path,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Return validation and provenance rows after checking added pair keys."""
    balanced = read_balanced_rows(balanced_path)
    reviewed = read_expert_rows(expert_path)
    seen = {(row["premise"], row["hypothesis"]) for row in balanced}
    for row in reviewed:
        pair = (row["premise"], row["hypothesis"])
        if pair in seen:
            raise ValueError(f"{expert_path}: duplicate pair in row {row['workbook_row']}")
        seen.add(pair)
    combined = balanced + [{field: row[field] for field in FIELDS} for row in reviewed]
    return combined, reviewed


def write_csv(path: Path, rows: list[dict[str, str]], fields: tuple[str, ...]) -> None:
    """Write rows as UTF-8 CSV in the specified column order."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    """Build the active validation CSV and its reviewed-row provenance."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--balanced", type=Path, default=DEFAULT_DIR / "val_balanced.csv")
    parser.add_argument("--expert", type=Path,
                        default=DEFAULT_DIR / "val_added_for_human_expert_done.xlsx")
    parser.add_argument("--output", type=Path, default=DEFAULT_DIR / "val.csv")
    parser.add_argument("--provenance", type=Path,
                        default=DEFAULT_DIR / "val_expert_provenance.csv")
    args = parser.parse_args()
    rows, reviewed = build_validation(args.balanced, args.expert)
    write_csv(args.output, rows, FIELDS)
    write_csv(args.provenance, reviewed, PROVENANCE_FIELDS)
    print(f"balanced={len(rows) - len(reviewed)} reviewed={len(reviewed)} total={len(rows)}")


if __name__ == "__main__":
    main()
