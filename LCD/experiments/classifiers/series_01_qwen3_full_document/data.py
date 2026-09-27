"""Copy expert-reviewed Full pairs with their complete source decisions."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from LCD.shared.autotest_scoring import normalize_subject_key
from LCD.shared.common import REPOSITORY_ROOT, file_sha256
from LCD.shared.documents import read_docx_text

DEFAULT_WORKBOOKS = REPOSITORY_ROOT / "local/data/benchmarks/autotest/Full"
DEFAULT_DOCUMENTS = REPOSITORY_ROOT / "local/data/benchmarks/test_docx/Full"
DEFAULT_OUTPUT = REPOSITORY_ROOT / "local/experiments/classifiers/series_01_qwen3_full_document"
LABELS = frozenset(("contradiction", "entailment", "not mentioned"))
COLUMNS = ("workbook", "workbook_row", "document", "hypothesis", "premise",
           "article_number", "expert_label", "document_context")


def match_sources(workbook_dir: Path, document_dir: Path) -> list[tuple[Path, Path]]:
    """Return unique reviewed-workbook/DOCX pairs, rejecting missing or extra files."""
    if not workbook_dir.is_dir() or not document_dir.is_dir():
        raise FileNotFoundError("Both Full workbook and DOCX directories are required")
    workbooks = sorted(p for p in workbook_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    documents = sorted(p for p in document_dir.glob("*.docx") if not p.name.startswith("~$"))
    if not workbooks or not documents:
        raise ValueError("Full workbook and DOCX directories must be nonempty")

    def keyed(paths: list[Path]) -> dict[str, Path]:
        result: dict[str, Path] = {}
        for path in paths:
            key = normalize_subject_key(path)
            if key in result:
                raise ValueError(f"Ambiguous subject {key!r}: {result[key].name}, {path.name}")
            result[key] = path
        return result

    xlsx = keyed(workbooks)
    docx = keyed(documents)
    if xlsx.keys() != docx.keys():
        raise ValueError(f"Unmatched Full sources: XLSX-only={sorted(xlsx.keys() - docx.keys())}; "
                         f"DOCX-only={sorted(docx.keys() - xlsx.keys())}")
    return [(xlsx[key], docx[key]) for key in sorted(xlsx)]


def read_reviewed_rows(workbook: Path, document: Path) -> list[dict[str, str | int]]:
    """Attach one DOCX text to each expert-labeled row, preserving Excel row numbers."""
    from openpyxl import load_workbook

    context = read_docx_text(document)
    if not context.strip():
        raise ValueError(f"Empty document: {document}")
    sheet = load_workbook(workbook, read_only=True, data_only=True).active
    values = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(values)]
    required = ("hypothesis", "premise", "article_number", "expert_label")
    if len(headers) != len(set(headers)) or any(name not in headers for name in required):
        raise ValueError(f"{workbook}: missing or duplicate required headers")
    positions = {name: headers.index(name) for name in required}
    rows: list[dict[str, str | int]] = []
    for excel_row, cells in enumerate(values, start=2):
        if all(value is None or str(value).strip() == "" for value in cells):
            continue
        fields = {name: str(cells[index]) if index < len(cells) and cells[index] is not None else ""
                  for name, index in positions.items()}
        if any(not fields[name].strip() for name in required[:3]):
            raise ValueError(f"{workbook.name}:{excel_row}: blank pair or citation")
        label = fields["expert_label"].strip().casefold()
        if label not in LABELS:
            raise ValueError(f"{workbook.name}:{excel_row}: invalid expert_label {fields['expert_label']!r}")
        rows.append({"workbook": workbook.name, "workbook_row": excel_row,
                     "document": document.name, "hypothesis": fields["hypothesis"],
                     "premise": fields["premise"], "article_number": fields["article_number"],
                     "expert_label": label, "document_context": context})
    if not rows:
        raise ValueError(f"No reviewed rows in {workbook}")
    return rows


def export_full(workbook_dir: Path = DEFAULT_WORKBOOKS,
                document_dir: Path = DEFAULT_DOCUMENTS,
                output_dir: Path = DEFAULT_OUTPUT) -> tuple[Path, Path]:
    """Validate all sources, then write the ignored training CSV and audit manifest."""
    pairs = match_sources(Path(workbook_dir), Path(document_dir))
    rows = [row for workbook, document in pairs for row in read_reviewed_rows(workbook, document)]
    manifest = {"source_role": "Full expert-reviewed training data", "pair_count": len(pairs),
                "row_count": len(rows), "labels": dict(sorted(Counter(r["expert_label"] for r in rows).items())),
                "sources": [{"workbook": workbook.name, "workbook_sha256": file_sha256(workbook),
                             "document": document.name, "document_sha256": file_sha256(document),
                             "rows": sum(r["workbook"] == workbook.name for r in rows)}
                            for workbook, document in pairs]}
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path, manifest_path = output_dir / "full_document_train.csv", output_dir / "manifest.json"
    with csv_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    manifest["csv_sha256"] = file_sha256(csv_path)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return csv_path, manifest_path


def main() -> None:
    """Run the local Full copy CLI."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbooks", type=Path, default=DEFAULT_WORKBOOKS)
    parser.add_argument("--documents", type=Path, default=DEFAULT_DOCUMENTS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(*export_full(args.workbooks, args.documents, args.output_dir), sep="\n")


if __name__ == "__main__":
    main()
