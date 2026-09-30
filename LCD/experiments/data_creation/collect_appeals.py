"""Collect verified Russian administrative-offence appeal decisions.

The collector searches SudAct, verifies the published body against the expert
criteria in ``instructions.md``, and stores the source HTML, extracted text,
DOCX copy, and provenance.  Importing this module has no side effects.

Example::

    python -m LCD.experiments.data_creation.collect_appeals \
        --output local/data/classification/dataset0929/appeals \
        --transport curl --target 5
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
from html.parser import HTMLParser
import json
import re
import shutil
import subprocess
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterable, Sequence
from urllib.parse import urlencode
from urllib.request import Request, urlopen


SUDACT_DOCUMENT = "https://sudact.ru/regular/doc/{document_id}/"
SUDACT_SEARCH = "https://sudact.ru/regular/doc_ajax/"
DEFAULT_OUTPUT = Path("local/data/classification/dataset0929/appeals")
MAIN_ARTICLES = (
    "18.8", "18.9", "18.10", "18.11", "18.12", "18.15", "18.16",
    "18.17", "18.18", "18.19", "18.20", "19.27",
)
DOMAIN_SHIFT_ARTICLES = ("20.20", "20.21", "6.9")
GROUPS = {"main": MAIN_ARTICLES, "domain_shift": DOMAIN_SHIFT_ARTICLES}
VERIFIED_SEEDS: dict[str, tuple[str, ...]] = {
    "18.19": ("WLaBX4Rzt4x", "qrO50ckIJoYc"),
    "18.20": (
        "Q8YbVS1XUjxv", "dJIEds3II2f", "AG93KZI9eGYX", "6wsTntzLOUWJ",
    ),
}
USER_AGENT = "Mozilla/5.0 (compatible; LegalContradictionDetection/1.0; research)"

PROVENANCE_FIELDS = (
    "document_id", "filename", "source_html", "source_text", "source_url",
    "retrieval_date", "retrieved_at_utc", "court", "case_number", "title",
    "article", "content_sha256", "evidence_chapter30",
    "evidence_challenged_ruling", "evidence_complaint", "evidence_article",
    "evidence_disposition", "disposition", "group", "source",
)
FAILURE_FIELDS = (
    "recorded_at_utc", "group", "article", "stage", "source_url", "reason",
)


class VisibleTextParser(HTMLParser):
    """Extract visible text while preserving useful paragraph boundaries."""

    BLOCK_TAGS = {"br", "p", "center", "div", "h1", "h2", "h3", "li"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.ignored_depth += 1
        elif not self.ignored_depth and tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.ignored_depth:
            self.ignored_depth -= 1
        elif not self.ignored_depth and tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth:
            self.parts.append(data)

    def paragraphs(self) -> list[str]:
        text = "".join(self.parts).replace("\xa0", " ")
        return [
            re.sub(r"[ \t]+", " ", line).strip()
            for line in text.splitlines()
            if line.strip()
        ]


@dataclass(frozen=True)
class PublishedDecision:
    """A parsed SudAct decision page."""

    document_id: str
    url: str
    title: str
    court: str
    case_number: str
    paragraphs: tuple[str, ...]
    source_html: str

    @property
    def text(self) -> str:
        return "\n".join(self.paragraphs)


@dataclass(frozen=True)
class Verification:
    """Evidence supporting acceptance of one appeal decision."""

    accepted: bool
    reason: str
    disposition: str = ""
    chapter30: str = ""
    challenged_ruling: str = ""
    complaint: str = ""
    article: str = ""
    disposition_evidence: str = ""


def _plain(fragment: str) -> str:
    return html.unescape(re.sub(r"<[^>]+>", "", fragment)).strip()


def parse_sudact_page(page: str, document_id: str, url: str) -> PublishedDecision | None:
    """Parse only the published decision body from a SudAct HTML page."""

    title_match = re.search(r"<h1[^>]*>(.*?)</h1>", page, re.I | re.S)
    court_match = re.search(r'<div class="b-justice"[^>]*>(.*?)</div>', page, re.I | re.S)
    boundary = re.search(r'<hr class="hr-h1"[^>]*>', page, re.I)
    if not title_match or not boundary:
        return None

    remainder = page[boundary.end():]
    end_markers = [
        marker for marker in (
            remainder.find("<!--AdFox START-->"),
            remainder.find('<div class="h-col2'),
            remainder.find('<div class="relates-block'),
        ) if marker >= 0
    ]
    fragment = remainder[:min(end_markers)] if end_markers else remainder
    # SudAct publishes redaction placeholders such as <адрес> as literal text.
    fragment = re.sub(r"<([А-Яа-яЁё][^<>]{0,80})>", r"&lt;\1&gt;", fragment)
    parser = VisibleTextParser()
    parser.feed(fragment)
    paragraphs = tuple(parser.paragraphs())
    if not paragraphs:
        return None

    title = _plain(title_match.group(1))
    court = _plain(court_match.group(1)) if court_match else ""
    case = re.search(r"(?:по делу\s*)?№\s*([^\n]+?)(?:\s+от\s+|$)", title, re.I)
    return PublishedDecision(
        document_id=document_id,
        url=url,
        title=title,
        court=re.sub(r"\s+", " ", court),
        case_number=case.group(1).strip() if case else "",
        paragraphs=paragraphs,
        source_html=page,
    )


def article_pattern(article: str) -> re.Pattern[str]:
    """Build a boundary-safe pattern for a КоАП article number."""

    major, minor = article.split(".")
    number = rf"{re.escape(major)}\s*[.,]\s*{re.escape(minor)}"
    return re.compile(
        rf"(?:ст(?:атья|атье|атьи|атью|атьёй|атьей)?\.?\s*{number}|"
        rf"(?<![\d.,]){number}\s*(?:ст(?:атья|атье|атьи|атью|атьёй|атьей)?\.?)?)"
        rf"(?!\s*\d)",
        re.I,
    )


def _normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("ё", "е").replace("Ё", "Е")).strip()


def _snippet(text: str, match: re.Match[str], radius: int = 280) -> str:
    start = max(0, match.start() - radius)
    end = min(len(text), match.end() + radius)
    prefix = "…" if start else ""
    suffix = "…" if end < len(text) else ""
    return prefix + text[start:end].strip() + suffix


def _find(text: str, pattern: str, flags: int = re.I | re.S) -> tuple[re.Match[str] | None, str]:
    match = re.search(pattern, text, flags)
    return match, _snippet(text, match) if match else ""


DISPOSITIONS: tuple[tuple[str, str], ...] = (
    (
        "revoked_and_sent_jurisdiction",
        r"постановлени\w*.{0,700}?отменить.{0,700}?"
        r"(?:направить|передать).{0,300}?(?:подведомственност|подсудност|компетенц)",
    ),
    (
        "revoked_and_remanded",
        r"постановлени\w*.{0,700}?отменить.{0,700}?"
        r"(?:дело|материал\w*).{0,300}?(?:возвратить|направить).{0,250}?новое рассмотрение",
    ),
    (
        "revoked_and_closed",
        r"постановлени\w*.{0,700}?отменить.{0,700}?"
        r"производств\w*.{0,250}?прекратить",
    ),
    (
        "changed",
        r"постановлени\w*.{0,700}?(?<!без )изменить(?:\s|[.,:])",
    ),
    (
        "upheld_denied",
        r"постановлени\w*.{0,900}?оставить\s+без\s+изменения.{0,500}?"
        r"жалоб\w*.{0,200}?без\s+удовлетворения",
    ),
)


def verify_appeal(decision: PublishedDecision, article: str) -> Verification:
    """Verify all expert criteria from the decision body, never its title."""

    text = _normalise(decision.text)
    if len(text) < 600:
        return Verification(False, "published body is too short")

    operative_markers = list(re.finditer(
        r"(?:Р\s*Е\s*Ш\s*И\s*Л(?:А)?|П\s*О\s*С\s*Т\s*А\s*Н\s*О\s*В\s*И\s*Л(?:А)?|"
        r"О\s*П\s*Р\s*Е\s*Д\s*Е\s*Л\s*И\s*Л(?:А)?)\s*:??",
        text,
        re.I,
    ))
    if not operative_markers:
        return Verification(False, "no final operative section")
    operative = text[operative_markers[-1].end():]

    chapter_match, chapter_evidence = _find(
        text,
        r"(?:глав\w*\s*30\s*(?:КоАП|Кодекс\w*.{0,80}?административн)|"
        r"(?:ст\.?(?:\s*ст\.?)?|стать\w*)\s*30\s*[.,]\s*[1-9]"
        r"(?:\s*[-–,]\s*(?:30\s*[.,]\s*)?[1-9])?\s*"
        r"(?:КоАП|Кодекс\w*\s+Российской\s+Федерации\s+об\s+административных\s+правонарушениях))",
    )
    if not chapter_match:
        return Verification(False, "Chapter 30 review is not established")

    complaint_match, complaint_evidence = _find(
        text[: min(len(text), 7000)],
        r"рассмотр\w*.{0,1400}?(?:по\s+)?жалоб\w*.{0,800}?"
        r"на\s+(?:вынесенн\w*.{0,180}?)?постановлени\w*",
    )
    if not complaint_match:
        return Verification(False, "complaint against the ruling is not established")
    complaint_route = complaint_match.group(0)
    if re.search(r"\bна\s+(?:решени|определени)\w*", complaint_route, re.I):
        return Verification(False, "current complaint challenges a prior review decision")

    target = article_pattern(article)
    opening = text[: min(len(text), 10000)]
    challenged_match, challenged_evidence = _find(
        opening,
        r"постановлени(?:ем|е|я)\s+.{0,1200}?(?:признан\w*\s+виновн\w*|"
        r"привлечен\w*\s+к\s+административной\s+ответственности|"
        r"назначен\w*\s+административн\w*\s+наказан\w*)",
    )
    if not challenged_match:
        return Verification(False, "challenged administrative-offence ruling is not established")

    established_marker = re.search(
        r"У\s*С\s*Т\s*А\s*Н\s*О\s*В\s*И\s*Л(?:А)?\s*:??", text, re.I,
    )
    preamble_end = established_marker.start() if established_marker else min(len(text), 3500)
    # The current complaint starts at complaint_match. Searching only from
    # there to УСТАНОВИЛ binds the article to the appealed ruling named in the
    # preamble. An article mentioned later as offence history cannot qualify.
    complaint_preamble = text[complaint_match.start():preamble_end]
    preamble_article = target.search(complaint_preamble)
    operative_article = target.search(operative)
    if preamble_article:
        article_evidence = _snippet(complaint_preamble, preamble_article)
    elif operative_article:
        article_evidence = _snippet(operative, operative_article)
    else:
        return Verification(
            False,
            f"article {article} is not tied to the current complaint or operative outcome",
        )

    disposition_name = ""
    disposition_match: re.Match[str] | None = None
    for name, pattern in DISPOSITIONS:
        match = re.search(pattern, operative, re.I | re.S)
        if match:
            disposition_name = name
            disposition_match = match
            break
    if not disposition_match:
        return Verification(False, "operative section has no allowed final disposition")

    return Verification(
        accepted=True,
        reason="accepted",
        disposition=disposition_name,
        chapter30=chapter_evidence,
        challenged_ruling=challenged_evidence,
        complaint=complaint_evidence,
        article=article_evidence,
        disposition_evidence=_snippet(operative, disposition_match, radius=180),
    )


class Fetcher:
    """HTTP text fetcher with urllib and curl transports."""

    def __init__(self, transport: str = "auto", timeout: int = 30) -> None:
        self.transport = transport
        self.timeout = timeout

    def __call__(self, url: str) -> str:
        errors: list[str] = []
        if self.transport in {"auto", "urllib"}:
            try:
                request = Request(url, headers={
                    "User-Agent": USER_AGENT,
                    "X-Requested-With": "XMLHttpRequest",
                })
                with urlopen(request, timeout=self.timeout) as response:
                    return response.read().decode("utf-8")
            except Exception as exc:  # pragma: no cover - environment dependent
                errors.append(f"urllib: {exc}")
                if self.transport == "urllib":
                    raise RuntimeError("; ".join(errors)) from exc
        if self.transport in {"auto", "curl"}:
            command = [
                "curl", "-sSL", "--fail", "--max-time", str(self.timeout),
                "-A", USER_AGENT, "-H", "X-Requested-With: XMLHttpRequest", url,
            ]
            completed = subprocess.run(command, capture_output=True, check=False)
            if completed.returncode == 0:
                return completed.stdout.decode("utf-8")
            errors.append(f"curl: {completed.stderr.decode('utf-8', 'replace').strip()}")
        raise RuntimeError("; ".join(errors) or "no HTTP transport selected")


def search_sudact(
    fetch: Callable[[str], str], query: str, page: int, *, polls: int = 12,
    poll_delay: float = 0.5, case_number: str = "",
) -> list[str]:
    """Return ordered unique document IDs from one SudAct search page."""

    parameters: dict[str, str | int] = {"regular-txt": query, "page": page}
    if case_number:
        parameters["regular-case_doc"] = case_number
    url = SUDACT_SEARCH + "?" + urlencode(parameters)
    payload: dict[str, object] = {}
    for attempt in range(polls):
        payload = json.loads(fetch(url))
        if "content" in payload:
            break
        if attempt + 1 < polls:
            time.sleep(poll_delay)
    content = str(payload.get("content", ""))
    found = re.findall(r'/regular/doc/([A-Za-z0-9]+)/', content)
    return list(dict.fromkeys(found))


def _safe_stem(article: str, decision: PublishedDecision) -> str:
    case = re.sub(r"[^0-9A-Za-zА-Яа-яЁё_-]+", "_", decision.case_number).strip("_")
    article_slug = article.replace(".", "_")
    return f"article_{article_slug}_{case or 'case'}_{decision.document_id}"


def write_docx(path: Path, decision: PublishedDecision) -> None:
    """Write a readable DOCX copy of the full published decision text."""

    from docx import Document

    document = Document()
    document.core_properties.title = decision.title
    document.core_properties.source = decision.url
    document.add_heading(decision.title, level=1)
    if decision.court:
        document.add_paragraph(decision.court)
    document.add_paragraph(decision.url)
    for paragraph in decision.paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def _read_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as file:
        return list(csv.DictReader(file))


def _write_rows(path: Path, fields: Sequence[str], rows: Iterable[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def revalidate_saved_rows(
    output: Path, rows: Sequence[dict[str, str]], failures: list[dict[str, str]],
) -> list[dict[str, str]]:
    """Recheck saved rows and quarantine anything rejected by current rules."""

    accepted: list[dict[str, str]] = []
    for row in rows:
        html_relative = row.get("source_html", "")
        html_path = output / html_relative
        if not html_relative or not html_path.is_file():
            accepted.append(row)
            continue
        page = html_path.read_text(encoding="utf-8")
        decision = parse_sudact_page(
            page,
            row.get("document_id", ""),
            row.get("source_url", ""),
        )
        verification = verify_appeal(decision, row.get("article", "")) if decision else None
        if verification and verification.accepted:
            accepted.append(row)
            continue

        group = row.get("group", "unknown")
        article = row.get("article", "unknown")
        quarantine = output / "rejected" / group / article
        quarantine.mkdir(parents=True, exist_ok=True)
        for field in ("filename", "source_html", "source_text"):
            relative = row.get(field, "")
            source = output / relative
            if source.is_file():
                destination = quarantine / source.name
                if not destination.exists():
                    shutil.move(str(source), str(destination))
        reason = verification.reason if verification else "saved source could not be parsed"
        failures.append({
            "recorded_at_utc": datetime.now(timezone.utc).isoformat(),
            "group": group,
            "article": article,
            "stage": "revalidation",
            "source_url": row.get("source_url", ""),
            "reason": reason,
        })
        print(f"QUARANTINED {group}/{article} {row.get('document_id', '')}: {reason}", flush=True)
    return accepted


def _relative(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


def persist_decision(
    output: Path, group: str, article: str, decision: PublishedDecision,
    verification: Verification,
) -> dict[str, str]:
    """Persist one accepted decision and return its provenance row."""

    article_dir = output / group / article
    article_dir.mkdir(parents=True, exist_ok=True)
    stem = _safe_stem(article, decision)
    html_path = article_dir / f"{stem}.html"
    text_path = article_dir / f"{stem}.txt"
    docx_path = article_dir / f"{stem}.docx"
    html_path.write_text(decision.source_html, encoding="utf-8")
    text_path.write_text(decision.text + "\n", encoding="utf-8")
    write_docx(docx_path, decision)

    now = datetime.now(timezone.utc)
    return {
        "document_id": decision.document_id,
        "filename": _relative(docx_path, output),
        "source_html": _relative(html_path, output),
        "source_text": _relative(text_path, output),
        "source_url": decision.url,
        "retrieval_date": now.date().isoformat(),
        "retrieved_at_utc": now.isoformat(),
        "court": decision.court,
        "case_number": decision.case_number,
        "title": decision.title,
        "article": article,
        "content_sha256": hashlib.sha256(decision.text.encode("utf-8")).hexdigest(),
        "evidence_chapter30": verification.chapter30,
        "evidence_challenged_ruling": verification.challenged_ruling,
        "evidence_complaint": verification.complaint,
        "evidence_article": verification.article,
        "evidence_disposition": verification.disposition_evidence,
        "disposition": verification.disposition,
        "group": group,
        "source": "sudact.ru",
    }


def _queries(article: str) -> tuple[str, ...]:
    return (
        f"{article} жалобу постановление",
        f"{article} постановление оставить без изменения жалобу",
        f"{article} постановление отменить жалобу",
        f"{article} постановление изменить жалобу",
    )


def _search_plans(article: str) -> tuple[tuple[str, str], ...]:
    """Prefer case-number category 12, used for direct Chapter 30 review."""

    focused = (
        (article, "12"),
        (f"{article} КоАП", "12"),
        (f"{article} жалобу постановление", "12"),
    )
    return focused + tuple((query, "") for query in _queries(article))


def collect(
    *, output: Path, groups: Sequence[str], articles: Sequence[str] | None,
    target: int, max_pages: int, fetch: Callable[[str], str], delay: float = 0.2,
) -> dict[tuple[str, str], int]:
    """Collect appeals, resuming from the output provenance manifest."""

    if target < 1:
        raise ValueError("target must be positive")
    output.mkdir(parents=True, exist_ok=True)
    manifest_path = output / "provenance.csv"
    failure_path = output / "failures.csv"
    rows = _read_rows(manifest_path)
    failures = _read_rows(failure_path)
    rows = revalidate_saved_rows(output, rows, failures)
    _write_rows(manifest_path, PROVENANCE_FIELDS, rows)
    _write_rows(failure_path, FAILURE_FIELDS, failures)
    seen_pairs = {(row.get("article", ""), row.get("document_id", "")) for row in rows}
    seen_hashes = {row.get("content_sha256", "") for row in rows if row.get("content_sha256")}
    counts: dict[tuple[str, str], int] = {}
    for row in rows:
        key = (row.get("group", ""), row.get("article", ""))
        counts[key] = counts.get(key, 0) + 1

    selected_articles = set(articles) if articles else None
    for group in groups:
        if group not in GROUPS:
            raise ValueError(f"unknown group: {group}")
        for article in GROUPS[group]:
            if selected_articles is not None and article not in selected_articles:
                continue
            key = (group, article)
            current = counts.get(key, 0)
            if current >= target:
                print(f"SKIP {group}/{article}: already has {current}", flush=True)
                continue

            attempted_ids: set[str] = set()
            seed_ids = VERIFIED_SEEDS.get(article, ())
            for document_id in seed_ids:
                if current >= target or (article, document_id) in seen_pairs:
                    continue
                attempted_ids.add(document_id)
                url = SUDACT_DOCUMENT.format(document_id=document_id)
                try:
                    page_html = fetch(url)
                    decision = parse_sudact_page(page_html, document_id, url)
                    verification = verify_appeal(decision, article) if decision else None
                    if verification and verification.accepted:
                        digest = hashlib.sha256(decision.text.encode("utf-8")).hexdigest()
                        if digest not in seen_hashes:
                            row = persist_decision(output, group, article, decision, verification)
                            rows.append(row)
                            seen_pairs.add((article, document_id))
                            seen_hashes.add(digest)
                            current += 1
                            counts[key] = current
                            _write_rows(manifest_path, PROVENANCE_FIELDS, rows)
                            print(
                                f"SAVED {group}/{article} {current}/{target}: "
                                f"{document_id} {verification.disposition}",
                                flush=True,
                            )
                    else:
                        failures.append({
                            "recorded_at_utc": datetime.now(timezone.utc).isoformat(),
                            "group": group, "article": article, "stage": "seed_rejected",
                            "source_url": url,
                            "reason": verification.reason if verification else "source could not be parsed",
                        })
                except Exception as exc:
                    failures.append({
                        "recorded_at_utc": datetime.now(timezone.utc).isoformat(),
                        "group": group, "article": article, "stage": "seed_document",
                        "source_url": url, "reason": str(exc),
                    })
            for query, case_number in _search_plans(article):
                if current >= target:
                    break
                for page in range(1, max_pages + 1):
                    if current >= target:
                        break
                    try:
                        identifiers = search_sudact(
                            fetch, query, page, case_number=case_number,
                        )
                    except Exception as exc:
                        failures.append({
                            "recorded_at_utc": datetime.now(timezone.utc).isoformat(),
                            "group": group, "article": article, "stage": "search",
                            "source_url": SUDACT_SEARCH, "reason": str(exc),
                        })
                        print(f"SEARCH ERROR {group}/{article} page {page}: {exc}", flush=True)
                        break
                    print(
                        f"SEARCH {group}/{article} case={case_number or '*'} "
                        f"page {page}: {len(identifiers)} candidates",
                        flush=True,
                    )
                    if not identifiers:
                        break
                    for document_id in identifiers:
                        if current >= target:
                            break
                        if document_id in attempted_ids or (article, document_id) in seen_pairs:
                            continue
                        attempted_ids.add(document_id)
                        url = SUDACT_DOCUMENT.format(document_id=document_id)
                        try:
                            page_html = fetch(url)
                            decision = parse_sudact_page(page_html, document_id, url)
                            if not decision:
                                continue
                            verification = verify_appeal(decision, article)
                            if not verification.accepted:
                                continue
                            digest = hashlib.sha256(decision.text.encode("utf-8")).hexdigest()
                            if digest in seen_hashes:
                                continue
                            row = persist_decision(output, group, article, decision, verification)
                            rows.append(row)
                            seen_pairs.add((article, document_id))
                            seen_hashes.add(digest)
                            current += 1
                            counts[key] = current
                            _write_rows(manifest_path, PROVENANCE_FIELDS, rows)
                            print(
                                f"SAVED {group}/{article} {current}/{target}: "
                                f"{document_id} {verification.disposition}",
                                flush=True,
                            )
                        except Exception as exc:
                            failures.append({
                                "recorded_at_utc": datetime.now(timezone.utc).isoformat(),
                                "group": group, "article": article, "stage": "document",
                                "source_url": url, "reason": str(exc),
                            })
                            print(f"DOCUMENT ERROR {url}: {exc}", flush=True)
                        if delay:
                            time.sleep(delay)
            counts[key] = current
            if current < target:
                failures.append({
                    "recorded_at_utc": datetime.now(timezone.utc).isoformat(),
                    "group": group, "article": article, "stage": "shortfall",
                    "source_url": "", "reason": f"collected {current} of target {target}",
                })
            _write_rows(failure_path, FAILURE_FIELDS, failures)
            print(f"ARTICLE DONE {group}/{article}: {current}/{target}", flush=True)

    _write_rows(manifest_path, PROVENANCE_FIELDS, rows)
    _write_rows(failure_path, FAILURE_FIELDS, failures)
    return counts


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--group", choices=("all", *GROUPS), default="all",
        help="article group to collect",
    )
    parser.add_argument(
        "--article", action="append",
        help="limit collection to an article; may be repeated",
    )
    parser.add_argument("--target", type=int, default=5)
    parser.add_argument("--max-pages", type=int, default=8)
    parser.add_argument("--delay", type=float, default=0.2)
    parser.add_argument("--timeout", type=int, default=30)
    parser.add_argument("--transport", choices=("auto", "urllib", "curl"), default="auto")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    """CLI entrypoint for resumable appeal collection."""

    args = parse_args(argv)
    groups = tuple(GROUPS) if args.group == "all" else (args.group,)
    valid = {article for group in groups for article in GROUPS[group]}
    if args.article:
        unknown = set(args.article) - valid
        if unknown:
            raise SystemExit(f"articles outside selected group: {', '.join(sorted(unknown))}")
    counts = collect(
        output=args.output,
        groups=groups,
        articles=args.article,
        target=args.target,
        max_pages=args.max_pages,
        fetch=Fetcher(args.transport, args.timeout),
        delay=max(0.0, args.delay),
    )
    for (group, article), count in counts.items():
        if group in groups and (not args.article or article in args.article):
            print(f"COUNT {group}/{article} {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
