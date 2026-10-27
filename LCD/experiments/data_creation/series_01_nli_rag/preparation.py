"""Prepare auditable hypotheses and expert quotas for the rebuilt datasets.

The filtering policy in this module is deliberately local to dataset creation.
It keeps the historical inference filter unchanged while making every keep/drop
decision inspectable. Importing the module performs no I/O.
"""

from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Iterable, Mapping

from LCD.shared.documents import OTD_ABBREVIATION, OTD_PLACEHOLDER, POSTANOVIL_PATTERN


_NUMERIC_ONLY = re.compile(r"(?:[\d\s.,;:/№-]+)[.!?]?\Z")
_MASK_ONLY = re.compile(
    r"(?:\s|<[^>]+>|\[[^]]+]|\([^)]*(?:данн|изъят|скрыт)[^)]*\)|[-–—.,;:№])+\Z",
    re.IGNORECASE,
)
_SIGNATURE = re.compile(
    r"(?:<[^>]+>\s*)*(?:(?:мировой|федеральный)\s+)?(?:судья|председательствующ(?:ий|ая))"
    r"(?:\s*[:—/-])?(?:\s*(?:\([^)]*\)|/[^/]*/|[А-ЯЁA-Z][\w.-]*:?)){0,12}\s*[.!]?\Z",
    re.IGNORECASE,
)
_BANK_TOKEN = re.compile(
    r"(?i)(?<!\w)(?:бик|инн|кпп|уин|кбк|октмо|огрн|л/с|р/с|кор\.?\s*сч(?:ет)?|"
    r"расч[её]тн(?:ый|ого)\s+сч[её]т|реквизит\w*|банковск\w*\s+реквизит|банк\s+получателя|"
    r"получатель\s+платежа|уфк)(?!\w)"
)
_BANK_BOILERPLATE = re.compile(
    r"(?i)^(?:реквизит|банковск\w*\s+реквизит|наименование\s+(?:платежа|получателя)|"
    r"взыскатель|получатель\b|банк\b|отделени\w*.*\bбанк|номер\s+сч[её]та|"
    r"№\s*сч[её]та|корр?\.?/?сч(?:[её]т)?|р/сч|окато\b|"
    r"казначейск\w*\s+сч[её]т|единый\s+казначейск\w*\s+сч[её]т|"
    r"код\s+бюджетн\w*\s+кл[а-яё]*сификац|идентификатор\b|(?:в\s+)?гу\s+банка\s+россии\b|"
    r"наименование\s+банка|банк\s+получател|уфк\b)"
)
_NON_LEGAL_BOILERPLATE = re.compile(
    r"(?i)^(?:административный\s+штраф|подлинный\s+документ\s+подшит|копия\s+верна)\s*[.!]?\Z"
)
_PAYMENT_TITLE = re.compile(r"(?i)^(?:наименование|назначение)\s+платежа\b")
_SIGNATURE_CONTINUATION = re.compile(
    r"(?i)^(?:городского|районного|областного|краевого|верховного)\s+суда\s*:\s*"
    r"(?:[А-ЯЁ]\s*\.\s*){1,3}[А-ЯЁ][а-яё-]+\s*[.!]?\Z"
)
_BANK_BLOCK_CONTINUATION = re.compile(
    r"(?i)^(?:\([^)]*\b(?:омвд|умвд|мвд|уфк|казначейств|банк|получател)\w*\b[^)]*\)|"
    r"(?:омвд|умвд|мвд|уфк|гу\s+(?:мвд|банка))\b[^.!?]*)\s*[.,;]?\Z"
)
_SUBSTANTIVE = re.compile(
    r"(?i)\b(?:признать|виновн|назначить|наказани|постановлени\w*\s+(?:может\s+быть\s+)?"
    r"обжал|обжалован|жалоб|(?:не\s+позднее|в\s+течени[еи]|на\s+следующ\w+\s+день)|"
    r"срок\w*\s+не\s+позднее|ответственност|выдворени|рассроч|"
    r"отсроч|вступлени\w*\s+в\s+законную\s+силу|контрол\w*\s+за\s+исполнени)\b"
)


def sentence_records(text: str) -> list[dict]:
    """Split each non-empty paragraph with razdel and retain source offsets.

    Paragraph boundaries are hard boundaries. This prevents a signature or bank
    line from becoming part of the preceding legal sentence. Offsets refer to
    the supplied ``text`` and the returned text is the exact source substring.
    """
    from razdel import sentenize

    records: list[dict] = []
    offset = 0
    for line in text.splitlines(keepends=True):
        paragraph = line.rstrip("\r\n")
        protected = paragraph.replace(OTD_ABBREVIATION, OTD_PLACEHOLDER)
        for sentence in sentenize(protected):
            start, end = sentence.start, sentence.stop
            while start < end and paragraph[start].isspace():
                start += 1
            while end > start and paragraph[end - 1].isspace():
                end -= 1
            if start == end:
                continue
            records.append(
                {
                    "index": len(records),
                    "text": paragraph[start:end],
                    "start": offset + start,
                    "end": offset + end,
                }
            )
        offset += len(line)
    # ``splitlines(keepends=True)`` returns no item for the empty string and
    # handles a final unterminated paragraph normally.
    return records


def _filter_reason(text: str) -> str | None:
    value = " ".join(text.split()).strip()
    if not value:
        return "empty"
    if ((_SIGNATURE.fullmatch(value) or _SIGNATURE_CONTINUATION.fullmatch(value))
            and not _SUBSTANTIVE.search(value)):
        return "judge_signature"
    if _MASK_ONLY.fullmatch(value):
        return "masked_only"
    if _NUMERIC_ONLY.fullmatch(value):
        return "numeric_only"
    if _NON_LEGAL_BOILERPLATE.fullmatch(value):
        return "non_legal_boilerplate"
    # A payment-purpose title remains bank metadata even when it repeats an
    # article number, ``КоАП`` or the word ``штраф``.
    if _PAYMENT_TITLE.search(value):
        return "bank_or_payment_details"
    # Identifiers are common in the adjudication of corporate defendants. A
    # bank marker is therefore insufficient when the same sentence states a
    # ruling, deadline, right of appeal, or other legal proposition.
    if (_BANK_TOKEN.search(value) or _BANK_BOILERPLATE.search(value)) and not _SUBSTANTIVE.search(value):
        return "bank_or_payment_details"
    if re.search(r"(?i)\bквитанци\w*\b", value) and not _SUBSTANTIVE.search(value):
        return "payment_receipt_instruction"
    return None


def extract_hypotheses(text: str) -> tuple[list[dict], list[dict]]:
    """Extract and audit hypotheses after the final operative-section marker.

    All offsets remain relative to the complete decision. Missing markers and
    sections with no retained hypothesis fail explicitly rather than silently
    producing an unusable document.
    """
    matches = list(POSTANOVIL_PATTERN.finditer(text))
    if not matches:
        raise ValueError("Court decision has no ПОСТАНОВИЛ section")
    section_start = matches[-1].end()
    section = text[section_start:]
    if not section.strip():
        raise ValueError("ПОСТАНОВИЛ section is empty")

    kept: list[dict] = []
    audit: list[dict] = []
    payment_block = False
    for record in sentence_records(section):
        item = {
            **record,
            "start": record["start"] + section_start,
            "end": record["end"] + section_start,
        }
        reason = _filter_reason(item["text"])
        if reason is None and payment_block and _BANK_BLOCK_CONTINUATION.fullmatch(
            " ".join(item["text"].split()).strip()
        ):
            reason = "bank_block_continuation"
        audited = {**item, "kept": reason is None, "reason": reason or "kept"}
        audit.append(audited)
        if reason is None:
            kept.append(item)
            # A retained legal proposition ends any preceding payment block.
            payment_block = False
        elif reason in {
            "bank_or_payment_details",
            "bank_block_continuation",
            "numeric_only",
            "masked_only",
        }:
            payment_block = True
        else:
            payment_block = False
    if not kept:
        raise ValueError("ПОСТАНОВИЛ section has no relevant hypotheses")
    return kept, audit


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _weight(value: object, *, row_number: int) -> float:
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*%", _text(value))
    if match is None:
        raise ValueError(f"Expert row {row_number} has no percentage weight")
    return float(match.group(1).replace(",", "."))


def read_expert_rules(path: str | Path) -> list[dict]:
    """Read only workbook rules marked exactly ``Да`` for the main collection."""
    from openpyxl import load_workbook

    source = Path(path)
    workbook = load_workbook(source, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        positions = {name: index for index, name in enumerate(headers) if name}

        def column(name: str) -> int:
            if name not in positions:
                raise ValueError(f"Expert workbook is missing column: {name}")
            return positions[name]

        name_col = column("Сигнал / ошибка")
        kind_col = column("Тип сигнала / ошибки")
        include_col = column("Добавляем в текущую выборку (октябрь 2026)")
        weight_col = column(
            "Процент документов от общего количества синтетических документов текущей основной и отдельной выборок"
        )
        comment_col = column("Комментарий")
        instruction_col = column("Инструкция для внесения синтетической ошибки")
        prompt_col = column("Промпт для поиска сигналов / ошибок")
        example_columns = [positions[name] for name in ("Пример 1", "Пример 2") if name in positions]
        mapped_columns = {name_col, kind_col, include_col, weight_col, comment_col,
                          instruction_col, prompt_col, *example_columns}

        rules = []
        for row_number, row in enumerate(rows, 2):
            inclusion = " ".join(_text(row[include_col]).split())
            if inclusion != "Да":
                continue
            name = _text(row[name_col])
            if not name:
                raise ValueError(f"Eligible expert row {row_number} has no name")
            rules.append(
                {
                    "rule_id": f"xlsx_row_{row_number}",
                    "row_number": row_number,
                    "name": name,
                    "kind": _text(row[kind_col]),
                    "weight_percent": _weight(row[weight_col], row_number=row_number),
                    "comment": _text(row[comment_col]),
                    "instruction": _text(row[instruction_col]),
                    "verification_prompt": _text(row[prompt_col]),
                    "examples": [_text(row[index]) for index in example_columns if _text(row[index])],
                    "additional_fields": {name: _text(row[index]) for name, index in positions.items()
                                          if index not in mapped_columns},
                }
            )
        if not rules:
            raise ValueError("Expert workbook contains no rows marked exactly Да")
        return rules
    finally:
        workbook.close()


def quota_targets(
    rules: Iterable[Mapping[str, object]], total: int, *, minimum: int = 0
) -> dict[str, int]:
    """Allocate a total by relative rule weights using largest remainders."""
    items = list(rules)
    if total < 0 or minimum < 0:
        raise ValueError("total and minimum must be non-negative")
    if not items:
        if total:
            raise ValueError("Cannot allocate a positive total without rules")
        return {}
    if total < minimum * len(items):
        raise ValueError("total is smaller than the requested per-rule minimum")
    identifiers = [str(rule["rule_id"]) for rule in items]
    if len(set(identifiers)) != len(identifiers):
        raise ValueError("rule_id values must be unique")
    weights = [float(rule["weight_percent"]) for rule in items]
    if any(not math.isfinite(weight) or weight < 0 for weight in weights) or not sum(weights):
        raise ValueError("Rule weights must be finite, non-negative, and not all zero")

    remaining = total - minimum * len(items)
    exact = [remaining * weight / sum(weights) for weight in weights]
    floors = [math.floor(value) for value in exact]
    result = {identifier: minimum + floor for identifier, floor in zip(identifiers, floors)}
    leftovers = remaining - sum(floors)
    order = sorted(range(len(items)), key=lambda index: (-(exact[index] - floors[index]), index))
    for index in order[:leftovers]:
        result[identifiers[index]] += 1
    return result
