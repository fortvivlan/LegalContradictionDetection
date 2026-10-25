"""Apply the expert document exclusions with evidence and reversible archiving.

The default scope is the main original-ruling pool. No network or models are
used. Analysis and moving files are separate CLI operations; imports do no I/O.
"""

from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
from typing import Sequence

from LCD.shared.documents import (
    extract_operative_section, read_docx_text, split_russian_sentences, textcheck,
)

DEFAULT_ROOT = Path("local/data/classification/dataset0929/decisions")
REASONS = {
    "direct_federal_law": "а: самостоятельная ссылка на ФЗ в разделе ПОСТАНОВИЛ",
    "entrepreneur_mention": "б: привлекаемое лицо — ИП / индивидуальный предприниматель",
    "split_penalty": "в: статья, штраф и назначенное выдворение в разных предложениях",
    "split_departure_explanation": "в: обязанность выехать отделена от ссылки на КоАП",
    "poor_spacing": "г: пропущенные пробелы в тексте постановления",
    "multiple_offences": "д: несколько составов / эпизодов в одном постановлении",
}
EXPERT_FILES = {
    "5-31_2026_HGoret1XwegL.docx": "split_penalty",
    "5-82_2026_PPEueqPWsdHU.docx": "split_penalty",
    "5-416_2025_OAuXAcEXzdxX.docx": "split_penalty",
    "5-443_2025_U6EcOBnxiUg.docx": "split_penalty",
    "5-210_2024_ebWvmZpQwr1L.docx": "split_departure_explanation",
    "5-253_2023_8BbSGQdwOLQQ.docx": "split_departure_explanation",
    "5-453_2021_H6MzReLM1Yl.docx": "poor_spacing",
    "5-104_2021_f9OINEGKfjZg.docx": "multiple_offences",
}
FEDERAL = re.compile(
    r"федеральн[а-яё]*\s+закон(?:а|ом|у|е|ы|ов|ам|ами|ах)?\b|(?<!\w)Ф\s*З(?!\w)", re.I,
)
ENTREPRENEUR = re.compile(r"\bИП\b|индивидуальн[а-яё]*\s+предпринимател[а-яё]*", re.I)
CODE = re.compile(r"КоАП|КРФ\s*об\s*АП|Кодекс[а-яё]*\s+(?:Российской\s+Федерации|РФ)\s+об\s+административных", re.I)
ARTICLE = re.compile(r"\bст(?:\.|ать[а-яё]*)\s*(\d+(?:\.\d+)+)", re.I)
ACTION = re.compile(r"\b(?:назначить|назначив|подвергнуть|подвергнув|наложить|наказать)\b", re.I)
ADVICE = re.compile(r"^\s*(?:Разъясн|Согласно|В соответствии|В силу|За\s|В случае|При\s|Сумм[ау]\s)", re.I)
DEPARTURE = re.compile(
    r"обязан[а-яё]*\s+(?:самостоятельно\s+)?выехать|\bобязать\b[\s\S]{0,1200}\bвыехать\b", re.I,
)
SPACING = re.compile(
    r"емунаказание|ейнаказание|рождениявиновн[а-яё]*|рублейс\b|размере\d+|"
    r"год[а-яё]*рождения|виновн[а-яё]*в\s+совершении|предусмотренногоч|"
    r"[а-яё]{3,}[А-ЯЁ][а-яё]{3,}|рублейбез\b"
)
MULTI_REVIEW = re.compile(
    r"объедин[а-яё]*[^\n]{0,200}(?:производств|дел[ао])|"
    r"(?<![\w.])(?:два|двух|три|тр[её]х|нескольк[а-яё]*|[2-9]\d*)\s+"
    r"(?:самостоятельных\s+|административных\s+)?(?:эпизод[а-яё]*|правонарушени[йя])|"
    r"по каждому (?:эпизод|правонаруш)|по\s+\d+\s+эпизод", re.I,
)


def _positive_expulsion(sentence: str) -> bool:
    """Recognize an imposed expulsion, excluding negation and generic warnings."""
    if ADVICE.search(sentence) or not re.search(r"выдвор", sentence, re.I):
        return False
    if re.search(r"без[^.;\n]{0,100}выдвор|выдвор[^.;\n]{0,150}не\s+(?:применять|назначать)|"
                 r"не\s+(?:применять|назначать)[^.;\n]{0,100}выдвор", sentence, re.I):
        return False
    return bool(
        re.search(r"\bвыдворить\b", sentence, re.I)
        or (ACTION.search(sentence) and re.search(
            r"\bс\s+(?:[а-яё()]+\s+){0,8}выдворени[ея]м|"
            r"\bи\s+административного\s+выдворения", sentence, re.I))
        or re.search(r"выдворение[^.;\n]{0,150}(?:назначить|осуществить|произвести)", sentence, re.I)
    )


def entrepreneur_party_evidence(text: str, operative: str) -> list[str]:
    """Find entrepreneur status attached to the accused, rather than a third party."""
    evidence = []
    for paragraph in operative.splitlines():
        guilt = re.search(r"\bвиновн[а-яё]*\b", paragraph, re.I)
        if not guilt and not re.match(r"\s*привлечь", paragraph, re.I):
            continue
        penalty = ACTION.search(paragraph)
        subject = paragraph[:penalty.start()] if penalty else paragraph
        for match in ENTREPRENEUR.finditer(subject):
            prefix = subject[max(0, match.start() - 100):match.start()]
            if re.search(r"\b(?:у|к|от)\s*$|работа[а-яё]*[^,;]{0,75}$|представител[а-яё]*\s*$", prefix, re.I):
                continue
            evidence.append(paragraph)
    opening = re.split(r"\bу\s*с\s*т\s*а\s*н\s*о\s*в\s*и\s*л(?:\s*а)?\b", text, maxsplit=1, flags=re.I)[0]
    for match in ENTREPRENEUR.finditer(opening):
        prefix = opening[max(0, match.start() - 240):match.start()]
        if re.search(r"в\s+отношении\s*:?\s*(?:ФИО\d+[^\n]{0,160})?$|"
                     r"о\s+привлечении[^\n]{0,180}ответственности[^\n]{0,100}$", prefix, re.I):
            evidence.append(opening[max(0, match.start() - 200):match.end() + 200])
    return list(dict.fromkeys(evidence))


def analyze_text(text: str, filename: str = "", ip_scope: str = "party") -> dict:
    """Return exclusion evidence and review candidates, without changing files.

    Sentence identity uses the existing razdel splitter. Paragraphs additionally
    provide context for legal citations and malformed sentence boundaries.
    Uncertain multi-offence and spacing matches are review flags, not exclusions.
    """
    operative = extract_operative_section(text)
    evidence: dict[str, list[str]] = {}
    protected: list[str] = []
    flags: dict[str, list[str]] = {}

    def add(reason: str, passage: str) -> None:
        if passage not in evidence.setdefault(reason, []):
            evidence[reason].append(passage)

    if not operative:
        flags["missing_operative"] = ["Нет конечного раздела ПОСТАНОВИЛ"]
        operative = ""
    sentences = split_russian_sentences(operative)
    paragraphs = operative.splitlines()
    # Federal-law references are checked before irrelevant-sentence filtering:
    # the experts explicitly exclude payment boilerplate containing named laws.
    for paragraph in paragraphs:
        if not FEDERAL.search(paragraph):
            continue
        for sentence in split_russian_sentences(paragraph):
            if not FEDERAL.search(sentence):
                continue
            # A concrete FZ provision remains direct even beside a KoAP citation.
            direct_article = re.search(
                r"\bст(?:\.|ать[а-яё]*)\s*\d+(?:\.\d+)*\s*"
                r"(?:Ф\s*З\b|Федеральн[а-яё]*\s+закон)", sentence, re.I,
            )
            embedded = CODE.search(sentence) and re.search(
                r"(?:специальн[а-яё]*\s+учрежден|срок[а-яё]*\s+содержани|"
                r"в\s+случаях,?\s+установленных\s+федеральным\s+законом)", sentence, re.I,
            )
            if embedded and not direct_article:
                protected.append(sentence)
            else:
                add("direct_federal_law", sentence)

    if ip_scope == "any":
        for match in ENTREPRENEUR.finditer(text):
            add("entrepreneur_mention", text[max(0, match.start() - 100):match.end() + 180])
    else:
        for passage in entrepreneur_party_evidence(text, operative):
            add("entrepreneur_mention", passage)

    guilt = [i for i, s in enumerate(sentences)
             if re.search(r"\bвиновн[а-яё]*\b", s, re.I)
             and any(5 <= int(a.split(".")[0]) <= 21 for a in ARTICLE.findall(s))
             and not ADVICE.search(s)]
    fine = [i for i, s in enumerate(sentences)
            if ACTION.search(s) and re.search(r"штраф", s, re.I) and not ADVICE.search(s)]
    expulsion = [i for i, s in enumerate(sentences) if _positive_expulsion(s)]
    if guilt and fine and expulsion and not set(guilt).intersection(fine, expulsion):
        for i in sorted(set(guilt + fine + expulsion)):
            add("split_penalty", sentences[i])
    for i, sentence in enumerate(sentences):
        if DEPARTURE.search(sentence) and not CODE.search(sentence):
            # Exclude analogous detached departure explanations when an
            # expulsion was actually imposed; generic hypothetical warnings
            # alone do not establish an imposed expulsion.
            previous = sentences[i - 1] if i else ""
            if expulsion or (re.search(r"32\.10", previous) and CODE.search(previous)):
                add("split_departure_explanation", previous + "\n" + sentence)
            else:
                flags.setdefault("uncited_departure", []).append(sentence)

    for i in guilt:
        sentence = sentences[i]
        guilt_word = re.search(r"\bвиновн[а-яё]*\b", sentence, re.I)
        charge = sentence[guilt_word.end():]
        penalty = ACTION.search(charge)
        if penalty:
            charge = charge[:penalty.start()]
        citations = []
        for match in ARTICLE.finditer(charge):
            if not 5 <= int(match.group(1).split(".")[0]) <= 21:
                continue
            prefix = charge[max(0, match.start() - 40):match.start()]
            parts = re.findall(r"(?:ч\.|част[а-яё]*)\s*(\d+(?:\.\d+)*)", prefix, re.I)
            citations.append((match.group(1), parts[-1] if parts else ""))
        if (len(set(citations)) > 1
                or re.search(r"совершении\s+(?:административных\s+)?правонарушени(?:ях|й)\b", charge, re.I)
                or re.search(r"част(?:ями|ей)\s+\d+\s*(?:,|и)\s*\d+", charge, re.I)):
            add("multiple_offences", sentence)
    if len(guilt) > 1:
        for i in guilt:
            add("multiple_offences", sentences[i])
    for match in MULTI_REVIEW.finditer(text):
        flags.setdefault("possible_multiple_offences", []).append(
            text[max(0, match.start() - 150):match.end() + 400]
        )
    for match in SPACING.finditer(operative):
        flags.setdefault("possible_poor_spacing", []).append(
            operative[max(0, match.start() - 80):match.end() + 140]
        )
    if filename in EXPERT_FILES:
        reason = EXPERT_FILES[filename]
        if reason not in evidence:
            add(reason, "Явное исключение экспертов.\n" + operative)
    return {
        "reasons": sorted(evidence), "evidence": evidence, "review_flags": flags,
        "protected_federal_law_contexts": protected,
        "operative": operative, "sentences": sentences,
        "filtered_sentence_indices": [i for i, s in enumerate(sentences) if textcheck(s)],
        "expert_named": filename in EXPERT_FILES,
    }


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_manifest(path: Path) -> tuple[list[str], list[dict]]:
    with path.open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        return list(reader.fieldnames or []), list(reader)


def _write_manifest(path: Path, fields: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _within(root: Path, relative: str) -> Path:
    path = (root / relative).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError(f"Path outside root: {relative}")
    return path


def create_review(root: Path, output: Path, groups: Sequence[str] = ("main",),
                  ip_scope: str = "party", overrides: dict | None = None) -> dict:
    """Scan authoritative DOCX files and write a hash-checked review plan.

    Overrides are keyed by manifest filename and contain optional add_reasons,
    remove_reasons, reviewed_flags and a mandatory explanatory note. Reports
    retain provenance and extraction text; this document review assigns no NLI
    predictions, hypotheses, premises or train/dev/test roles.
    """
    root, output = root.resolve(), output.resolve()
    if output.is_relative_to(root):
        raise ValueError("Review output must be outside the active decisions tree")
    if (output / "plan.json").exists():
        raise FileExistsError(f"Review already exists: {output}")
    fields, rows = _read_manifest(root / "provenance.csv")
    sources = {row["filename"] for row in rows if row["group"] in groups}
    actual = {p.relative_to(root).as_posix() for group in groups
              for p in (root / group).rglob("*.docx")}
    if sources != actual:
        raise ValueError(f"Manifest/DOCX mismatch: {sources ^ actual}")
    overrides = overrides or {}
    if set(overrides) - sources:
        raise ValueError(f"Unknown override sources: {set(overrides) - sources}")
    records = []
    for row in rows:
        if row["group"] not in groups:
            continue
        source = row["filename"]
        path = _within(root, source)
        result = analyze_text(read_docx_text(path), path.name, ip_scope)
        override = overrides.get(source)
        if override:
            if not override.get("note"):
                raise ValueError(f"Missing override note: {source}")
            reasons = set(result["reasons"])
            reasons.update(override.get("add_reasons", []))
            reasons.difference_update(override.get("remove_reasons", []))
            if reasons - REASONS.keys():
                raise ValueError(f"Unknown reasons: {reasons - REASONS.keys()}")
            if result["expert_named"] and EXPERT_FILES[path.name] not in reasons:
                raise ValueError(f"Cannot undo explicit expert exclusion: {source}")
            result["reasons"] = sorted(reasons)
            for reason in override.get("add_reasons", []):
                result["evidence"].setdefault(reason, []).extend(
                    override.get("evidence", {}).get(reason, [override["note"]])
                )
        reviewed = set((override or {}).get("reviewed_flags", []))
        result["pending_review_flags"] = sorted(set(result["review_flags"]) - reviewed)
        records.append({
            "source": source, "provenance": row, **result,
            "action": "exclude" if result["reasons"] else "keep",
            "manual_review": override,
            "files": {field: {"path": row[field], "sha256": _digest(_within(root, row[field]))}
                      for field in ("filename", "source_text", "source_html")},
        })
    plan = {
        "root": str(root), "output": str(output),
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "groups": list(groups), "ip_scope": ip_scope,
        "manifest_sha256": _digest(root / "provenance.csv"),
        "manifest_fields": fields, "records": records,
    }
    output.mkdir(parents=True, exist_ok=True)
    (output / "plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    write_reports(plan, output)
    return plan


def write_reports(plan: dict, output: Path) -> dict:
    """Export a readable Russian report and an Excel review workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    records = plan["records"]
    summary = {
        "reviewed": len(records),
        "excluded": sum(r["action"] == "exclude" for r in records),
        "kept": sum(r["action"] == "keep" for r in records),
        "reason_counts_overlapping": dict(Counter(reason for r in records for reason in r["reasons"])),
        "pending_review_documents": sum(bool(r["pending_review_flags"]) for r in records),
        "by_article": {},
    }
    for r in records:
        article = r["provenance"]["article"]
        summary["by_article"].setdefault(article, {"keep": 0, "exclude": 0})[r["action"]] += 1
    (output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    ip_label = "только ИП как привлекаемое лицо" if plan["ip_scope"] == "party" else "любое упоминание ИП"
    lines = ["Проверка постановлений по критериям экспертов", "",
             f"Проверено: {summary['reviewed']}; исключено: {summary['excluded']}; оставлено: {summary['kept']}.",
             f"Критерий ИП: {ip_label}. Группы: {', '.join(plan['groups'])}.",
             "Числа по причинам пересекаются. Исключённые документы сохраняются в excluded/.", ""]
    for reason, count in summary["reason_counts_overlapping"].items():
        lines.append(f"{REASONS[reason]}: {count}")
    workbook = Workbook()
    totals = workbook.active
    totals.title = "Сводка"
    totals.append(["Показатель / причина", "Количество / значение"])
    for key, label in (("reviewed", "Проверено"), ("excluded", "Исключено"), ("kept", "Оставлено"),
                       ("pending_review_documents", "Непроверенные кандидаты")):
        totals.append([label, summary[key]])
    totals.append(["Критерий ИП", ip_label])
    for reason, count in summary["reason_counts_overlapping"].items():
        totals.append([REASONS[reason], count])
    totals.append(["Числа по причинам пересекаются", "Один документ может иметь несколько причин"])
    totals.column_dimensions["A"].width = 90
    totals.column_dimensions["B"].width = 65
    for title, subset in (("Исключены", [r for r in records if r["action"] == "exclude"]),
                          ("Оставлены", [r for r in records if r["action"] == "keep"])):
        sheet = workbook.create_sheet(title)
        sheet.append(["Документ", "Статья", "Причины", "Фрагменты", "Ручная проверка", "Ссылка на источник"])
        for r in subset:
            passages = "\n\n".join(k + ":\n" + "\n".join(v) for k, v in r["evidence"].items())
            sheet.append([r["source"], r["provenance"]["article"],
                          "\n".join(REASONS[x] for x in r["reasons"]), passages[:32000],
                          (r["manual_review"] or {}).get("note", ""), r["provenance"]["source_url"]])
            sheet.cell(sheet.max_row, 6).hyperlink = r["provenance"]["source_url"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCDEF", (65, 12, 60, 110, 90, 55)):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(output / "review.xlsx")
    for r in records:
        if r["action"] == "exclude" or r["manual_review"] or r["protected_federal_law_contexts"]:
            lines.extend(["", f"{r['action'].upper()} {r['source']}",
                          "; ".join(REASONS[x] for x in r["reasons"]),
                          (r["manual_review"] or {}).get("note", "")])
            for reason, passages in r["evidence"].items():
                if reason in r["reasons"]:
                    lines.extend([REASONS[reason], *passages])
            if r["protected_federal_law_contexts"]:
                lines.extend(["Допустимые вложенные ссылки на ФЗ:", *r["protected_federal_law_contexts"]])
    (output / "report.txt").write_text("\n".join(lines), encoding="utf-8")
    return summary


def apply_review(plan_path: Path) -> dict:
    """Archive excluded bundles and update active provenance, with rollback.

    Preflight checks every source hash, the manifest and destination collisions.
    The original manifest and excluded provenance are retained with the archive.
    A successful plan cannot be applied twice.
    """
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    root, output = Path(plan["root"]), Path(plan["output"])
    if (output / "applied.json").exists():
        raise FileExistsError("Review already applied")
    if root == output or output.is_relative_to(root):
        raise ValueError("Archive must be outside the active decisions tree")
    if _digest(root / "provenance.csv") != plan["manifest_sha256"]:
        raise ValueError("Active provenance changed since review")
    if any(r["pending_review_flags"] for r in plan["records"]):
        raise ValueError("Resolve pending review flags before applying")
    moves = []
    for record in plan["records"]:
        for item in record["files"].values():
            source = _within(root, item["path"])
            if _digest(source) != item["sha256"]:
                raise ValueError(f"Source changed since review: {source}")
            if record["action"] == "exclude":
                destination = _within(output / "excluded", item["path"])
                if destination.exists():
                    raise FileExistsError(destination)
                moves.append((source, destination))
    fields, rows = _read_manifest(root / "provenance.csv")
    rejected = {r["source"] for r in plan["records"] if r["action"] == "exclude"}
    backup = output / "before_review" / "provenance.csv"
    backup.parent.mkdir(parents=True, exist_ok=True)
    if backup.exists():
        if _digest(backup) != plan["manifest_sha256"]:
            raise ValueError("Conflicting backup manifest")
    else:
        shutil.copy2(root / "provenance.csv", backup)
    (output / "excluded").mkdir(exist_ok=True)
    _write_manifest(output / "excluded" / "provenance.csv", fields,
                    [r for r in rows if r["filename"] in rejected])
    completed = []
    try:
        for source, destination in moves:
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(source), str(destination))
            completed.append((source, destination))
        staged = root / "provenance.review.tmp"
        _write_manifest(staged, fields, [r for r in rows if r["filename"] not in rejected])
        staged.replace(root / "provenance.csv")
        result = {"excluded": len(rejected), "remaining_manifest_rows": len(rows) - len(rejected),
                  "moved_files": len(completed), "applied_at_utc": datetime.now(timezone.utc).isoformat()}
        (output / "applied.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    except Exception:
        for source, destination in reversed(completed):
            shutil.move(str(destination), str(source))
        shutil.copy2(backup, root / "provenance.csv")
        raise
    return result


def main(argv: Sequence[str] | None = None) -> int:
    """Create a review, or explicitly apply an existing reviewed plan."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--groups", nargs="+", choices=("main", "domain_shift"), default=["main"])
    parser.add_argument("--ip-scope", choices=("any", "party"), default="party")
    parser.add_argument("--overrides", type=Path)
    parser.add_argument("--apply", type=Path, metavar="PLAN_JSON")
    args = parser.parse_args(argv)
    if args.apply:
        print(json.dumps(apply_review(args.apply), ensure_ascii=False, indent=2))
    else:
        if not args.output:
            parser.error("--output is required when creating a review")
        overrides = json.loads(args.overrides.read_text(encoding="utf-8")) if args.overrides else None
        plan = create_review(args.root, args.output, args.groups, args.ip_scope, overrides)
        print(json.dumps({"reviewed": len(plan["records"]), "plan": str(args.output / "plan.json")}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
