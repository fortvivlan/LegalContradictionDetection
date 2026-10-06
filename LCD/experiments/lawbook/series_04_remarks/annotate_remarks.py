"""Extract DOCX remarks and export reviewable scope and entity annotations.

No actions occur on import. DOCX styles distinguish publisher commentary from
legal text. Numbered remarks become separate records; internal ``1)`` lists and
continuation paragraphs stay within their remark. Annotations are provisional.
"""

from __future__ import annotations

import argparse
from collections import Counter
import csv
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MARKER = re.compile(r"^Примечани(?P<form>е|я)\s*[.:]\s*(?P<body>.*)$", re.I)
NUMBER = re.compile(r"^(\d+(?:\.\d+)*)\.\s+")
ARTICLE = re.compile(r"^Статья\s+(\d+(?:\.\d+)*(?:\s+\d+)*)\.(?:\s|$)")
EDITORIAL = {"Комментарий", "Информация о версии", "Информация об изменениях",
             "Подзаголовок для информации об изменениях"}
NUMERALS = r"\d+(?:\.\d+)*(?:\s*(?:[-–—]|,|или|и)\s*\d+(?:\.\d+)*)*"
LOCAL_PARTS = re.compile(r"част(?:ь|и|ью|ей|ями|ях)\s+(" + NUMERALS + r")\s+настоящей\s+стать[а-я]+", re.I)
CATEGORIES = {
    0: ("КоАП в целом", "Общее определение или правило для всего КоАП. Применяется только там, где встречается соответствующее понятие; это не означает действие каждого положения примечания на каждую статью."),
    1: ("Одна статья", "Правило или определение для одной статьи без явно выделенного структурного элемента. Тематические ограничения, круг лиц и условия сохраняются в тексте и поле scope_qualifier."),
    2: ("Элементы одной статьи", "Одна или несколько частей, абзацев или пунктов одной статьи. Номера частей раскрыты по одному объекту связи; номер примечания не является номером части."),
    3: ("Несколько статей", "Правило охватывает несколько статей либо их отдельные части. Разные уровни детализации внутри перечня статей допустимы: отдельная категория смешанного охвата для этого не нужна."),
    4: ("Глава или раздел", "Общее правило для главы или раздела КоАП. Исключения внутри главы сохраняются как условия; глава не разворачивается автоматически в перечень всех статей."),
    5: ("Смешанный охват", "В одном неразделимом примечании сочетаются общее правило для КоАП или главы и отдельные правила для указанных статей/частей; либо для статьи и остальной главы установлены разные определения. Каждая область получает отдельную связь и ограничение."),
    6: ("Другое примечание", "Непосредственно уточняет или ограничивает другое примечание. Цель связи — примечание с номером и статьей; его действие на основные нормы можно получить позднее через эту связь, без подмены номера примечания номером части."),
    7: ("Нормы вне КоАП", "Прямо распространяет правило на нормы вне КоАП, например на определенные правонарушения по законам субъектов РФ. Указывается группа норм; точные региональные статьи из текста установить нельзя. Простая ссылка на внешний закон как источник определения не образует эту категорию."),
    8: ("Тематическая группа; границы уточнить", "Область задана описанием правонарушений без достаточного перечня структурных единиц. Сохраняются тематическая сущность и возможные статьи для проверки; окончательные ссылки не выдумываются."),
    9: ("Утратившее силу", "В документе осталась только отметка об утрате силы. Сохраняется место размещения; текущая действующая область применения не назначается. Такой объект нельзя подключать как действующее правило."),
}


def entity(level: str, *, article: str = "", part: str = "", paragraph: str = "",
           chapter: str = "", section: str = "", remark_number: str = "",
           code_id: str = "koap_rf", qualifier: str = "") -> dict:
    """Build a structured entity and deterministic, edition-independent link ID."""
    values = {"code_id": code_id, "level": level, "article": article, "part": part,
              "paragraph": paragraph, "chapter": chapter, "section": section,
              "remark_number": remark_number, "scope_qualifier": qualifier}
    identity = code_id
    for name in ("section", "chapter", "article", "part", "paragraph", "remark_number"):
        if values[name]:
            identity += f":{'remark' if name == 'remark_number' else name}:{values[name]}"
    if level == "norm_group":
        identity += ":norm_group:" + hashlib.sha256(qualifier.encode()).hexdigest()[:12]
    values["entity_id"] = identity
    return values


def entity_label(target: dict) -> str:
    """Format a machine-readable entity for a Russian expert workbook."""
    label = "КоАП РФ" if target["code_id"] == "koap_rf" else "Законы субъектов РФ"
    for name, word in (("section", "раздел"), ("chapter", "глава"), ("article", "статья"),
                       ("part", "часть"), ("paragraph", "абзац"), ("remark_number", "примечание")):
        if target.get(name):
            label += f", {word} {target[name]}"
    if target["level"] == "code":
        label += " в целом"
    if target.get("scope_qualifier"):
        label += " [" + target["scope_qualifier"] + "]"
    return label


def expand_numbers(text: str) -> list[str]:
    """Expand ascending ranges with a common prefix; fail on ambiguous ranges."""
    tokens = re.findall(r"\d+(?:\.\d+)*|[-–—]", text)
    result = []
    i = 0
    while i < len(tokens):
        first = tokens[i]
        if i + 1 < len(tokens) and tokens[i + 1] in {"-", "–", "—"}:
            last = tokens[i + 2]
            a, b = first.split("."), last.split(".")
            if a[:-1] != b[:-1] or int(b[-1]) < int(a[-1]) or int(b[-1]) - int(a[-1]) > 100:
                raise ValueError(f"Ambiguous range requiring annotation: {first} - {last}")
            prefix = ".".join(a[:-1])
            result.extend((prefix + "." if prefix else "") + str(n)
                          for n in range(int(a[-1]), int(b[-1]) + 1))
            i += 3
        else:
            result.append(first)
            i += 1
    return list(dict.fromkeys(result))


def suggest_annotation(record: dict) -> dict:
    """Suggest conservative scopes; edition-specific reviewed overrides win.

    These rules deliberately do not resolve every citation as an application
    target. Complex lists and indirect scopes need curated local annotations.
    """
    text, article = record["remark_text"], record["article"]
    annotation = {"category": 1, "targets": [entity("article", article=article)],
                  "annotation_method": "contextual_article", "review_priority": "обычный",
                  "annotation_reason": "Область восстановлена по статье размещения; текст не задает более широкую область.",
                  "candidate_targets": [], "review_status": "предварительно; эксперт не проверял"}
    if re.search(r"\bутратил[оиа]\s+силу|\bутратили\s+силу", text, re.I):
        annotation.update(category=9, targets=[], annotation_method="repealed_placeholder",
                          annotation_reason="Сохранена отметка об утрате силы; действующие связи отсутствуют.")
    elif re.search(r"(?:для целей|в)\s+настоящ(?:его|ем)\s+Кодекс", text, re.I):
        annotation.update(category=0, targets=[entity("code")], annotation_method="explicit_text",
                          annotation_reason="В тексте прямо указан настоящий Кодекс.")
    elif re.search(r"настоящ(?:ей|ую)\s+глав", text, re.I):
        chapter = re.match(r"Глава\s+(\d+)", record["chapter"])
        if not chapter:
            raise ValueError(f"No chapter context for {record['remark_id']}")
        annotation.update(category=4, targets=[entity("chapter", chapter=chapter.group(1))],
                          annotation_method="explicit_text", annotation_reason="Прямо указана настоящая глава.")
    else:
        matches = list(LOCAL_PARTS.finditer(text))
        if matches:
            parts = list(dict.fromkeys(p for match in matches for p in expand_numbers(match.group(1))))
            annotation.update(category=2, targets=[entity("part", article=article, part=p) for p in parts],
                              annotation_method="explicit_text", annotation_reason="Прямо указаны части настоящей статьи.")
        elif re.search(r"настоящ(?:ей|ую|ая)\s+стать", text, re.I):
            annotation.update(annotation_method="explicit_text",
                              annotation_reason="Прямо указана настоящая статья.")
    return annotation


def read_paragraphs(source: Path) -> list[dict]:
    """Read DOCX paragraphs, including tables, in XML order with style names."""
    with ZipFile(source) as archive:
        styles = ET.fromstring(archive.read("word/styles.xml"))
        names = {node.get(W + "styleId"): node.find(W + "name").get(W + "val")
                 for node in styles.findall(W + "style") if node.find(W + "name") is not None}
        body = ET.fromstring(archive.read("word/document.xml")).find(W + "body")
    result = []
    for index, node in enumerate(body.iter(W + "p"), 1):
        text = "".join(n.text or "" if n.tag == W + "t" else " "
                       for n in node.iter() if n.tag in {W + "t", W + "tab", W + "br"})
        style = node.find(W + "pPr/" + W + "pStyle")
        result.append({"index": index, "text": re.sub(r"\s+", " ", text).strip(),
                       "style": names.get(style.get(W + "val"), "") if style is not None else "Normal"})
    return result


def extract_remarks(paragraphs: list[dict]) -> list[dict]:
    """Extract individual remarks with article, block and paragraph provenance.

    Keep repealed placeholders as records for explicit review. Skip editorial
    paragraphs even when they interrupt a numbered list. Indices are one-based
    XML paragraph positions and include empty paragraphs and table paragraphs.
    """
    records = []
    article = title = chapter = section = ""
    active = None
    block = 0
    plural = False
    pending = False

    def flush():
        nonlocal active
        if active:
            active["remark_text"] = "\n".join(active.pop("texts"))
            active["remark_id"] = f"remark_{len(records) + 1:03d}"
            records.append(active)
            active = None

    def start(text, index, number=""):
        nonlocal active
        active = {"article": article, "article_title": title, "chapter": chapter,
                  "section": section, "remark_number": number, "block_index": block,
                  "paragraph_start": index, "paragraph_end": index, "texts": [text]}

    for paragraph in paragraphs:
        text, style, index = paragraph["text"], paragraph["style"], paragraph["index"]
        if style in EDITORIAL or not text:
            continue
        heading = ARTICLE.match(text)
        if heading:
            flush()
            article = re.sub(r"\s+", ".", heading.group(1))
            title = text
            pending = plural = False
            continue
        if re.match(r"^(?:Глава|Раздел)\s+", text):
            flush()
            pending = plural = False
            if text.startswith("Глава"):
                chapter = text
            else:
                section, chapter = text, ""
            continue
        marker = MARKER.match(text)
        if marker:
            if not article:
                raise ValueError(f"Remark outside an article at paragraph {index}")
            flush()
            block += 1
            plural = marker.group("form").lower() == "я"
            pending = True
            body = marker.group("body")
            if body:
                numbered = NUMBER.match(body) if plural else None
                start(body if numbered else text, index, numbered.group(1) if numbered else "")
            continue
        if pending:
            numbered = NUMBER.match(text) if plural else None
            if numbered:
                flush()
                start(text, index, numbered.group(1))
            elif active:
                active["texts"].append(text)
                active["paragraph_end"] = index
            else:
                raise ValueError(f"Expected numbered remark at paragraph {index}: {text}")
    flush()
    return records


def annotate_records(records: list[dict], overrides: dict | None = None) -> list[dict]:
    """Apply local text-hash-checked annotations to extracted records.

    Overrides are keyed by remark_id and contain a text_sha256 plus annotation
    fields. A changed remark fails rather than silently reusing stale links.
    """
    overrides = overrides or {}
    unknown = set(overrides) - {r["remark_id"] for r in records}
    if unknown:
        raise ValueError(f"Unknown annotation IDs: {sorted(unknown)}")
    result = []
    for record in records:
        digest = hashlib.sha256(record["remark_text"].encode("utf-8")).hexdigest()
        override = dict(overrides.get(record["remark_id"], {}))
        if override and override.pop("text_sha256", None) != digest:
            raise ValueError(f"Stale annotation text for {record['remark_id']}")
        annotation = {**suggest_annotation(record), **override}
        if annotation["category"] not in CATEGORIES:
            raise ValueError(f"Unknown category for {record['remark_id']}")
        if annotation["category"] == 9 and annotation["targets"]:
            raise ValueError("Repealed placeholders cannot have active application links")
        if annotation["category"] != 9 and not annotation["targets"]:
            raise ValueError(f"Missing entity for {record['remark_id']}")
        result.append({**record, **annotation, "text_sha256": digest})
    return result


def write_categories(path: Path, counts: Counter) -> None:
    """Write the Russian taxonomy and annotation/linking conventions as UTF-8."""
    lines = ["Категории области применения примечаний КоАП РФ", "",
             "Разметка предварительная, выполнена по предоставленному DOCX. Эксперт еще не проверял ее.",
             "Одна строка — одно ненумерованное примечание или один нумерованный элемент блока «Примечания».",
             "Внутренние списки вида 1), 2) и продолжения абзацами остаются частью той же строки.", ""]
    for number, (name, description) in CATEGORIES.items():
        lines += [f"{number} — {name} (строк: {counts[number]})", description, ""]
    lines += ["Правила аннотирования", "",
              "Категория описывает непосредственную область применения, а не все цитаты в тексте.",
              "Ссылки на условия, исключения, источники определений и прежние нарушения не превращаются автоматически в цели связи.",
              "Пример: примечание к статье 1.5 относится к части 3 статьи 1.5; глава 12 и перечисленные иные статьи описывают исключения.",
              "Если в тексте прямо указана вся статья, не сужаем ее до части только по встречаемости термина.",
              "Если номера в тексте отсутствуют, отдельные части иногда восстановлены по диспозициям статьи; такие выводы помечены для приоритетной проверки.",
              "Исключения, круг лиц, предмет и другие условия остаются в полном тексте и при необходимости в scope_qualifier.",
              "Пустой targets_json у категории 9 означает отсутствие действующей области, а не потерю статьи размещения.",
              "Связи с другими примечаниями сохраняются непосредственно; транзитивное распространение на статьи здесь не вычисляется.", "",
              "Сущности для последующего связывания", "",
              "В CSV targets_json содержит массив целей, а remark_targets.csv — одну связь на строку.",
              "code_id, level, article, part, paragraph, chapter, section, remark_number задают структурную единицу.",
              "entity_id — стабильный структурный идентификатор, например koap_rf:article:1.5:part:3.",
              "Для другого примечания: koap_rf:article:16.2:remark:2. remark_number не является part.",
              "level=norm_group обозначает тематическую группу без исчерпывающих номеров; идентификатор включает хеш ее описания.",
              "Сущности edition-independent: при связывании всегда добавляйте source_sha256/редакцию, чтобы не смешать редакции кодекса.",
              "remark_id стабилен в пределах данного извлечения, а source_sha256 и text_sha256 защищают от применения устаревшей разметки.",
              "candidate_targets_json содержит предложения для эксперта; они не входят в действующие связи remark_targets.csv.",
              "Для снятых примечаний location_entity_id показывает только место размещения.",
              "Дата редакции берется из имени предоставленного файла; действие норм на конкретную дату отдельно не проверялось.", "",
              "Проверка XLSX", "",
              "Первые три колонки: полный текст, категория, сущности применения. Фильтры и закрепленная строка включены.",
              "Колонки «Категория эксперта», «Сущности эксперта» и «Комментарий эксперта» оставлены для исправлений.",
              "Длинный текст хранится целиком в ячейке; если он не помещается по высоте, откройте ячейку или разверните строку формул.",
              "Листы «Связи» и «Категории» содержат структуру целей и справочник. Лист «Сведения» фиксирует источник и метод."]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    """Write a UTF-8-with-BOM CSV with full multiline text and stable headers."""
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def export_records(records: list[dict], source: Path, output_dir: Path) -> dict:
    """Write CSV, taxonomy, editable expert XLSX and reproducibility metadata."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    output_dir.mkdir(parents=True, exist_ok=True)
    rows, links = [], []
    for record in records:
        category = record["category"]
        location = entity("article", article=record["article"])
        row = {k: record[k] for k in ("remark_id", "article", "article_title", "chapter", "section",
               "remark_number", "block_index", "paragraph_start", "paragraph_end", "remark_text",
               "annotation_method", "annotation_reason", "review_priority", "review_status", "text_sha256")}
        row.update(source_document=str(source), source_sha256=source_hash,
                   location_entity_id=location["entity_id"], category_id=category,
                   category_label=f"{category} — {CATEGORIES[category][0]}",
                   applies_to="; ".join(entity_label(t) for t in record["targets"]) or "Нет действующей области; см. место размещения",
                   targets_json=json.dumps(record["targets"], ensure_ascii=False),
                   candidate_targets_json=json.dumps(record["candidate_targets"], ensure_ascii=False))
        rows.append(row)
        for index, target in enumerate(record["targets"], 1):
            links.append({"remark_id": record["remark_id"], "source_sha256": source_hash,
                          "target_index": index, "category_id": category, **target,
                          "annotation_method": record["annotation_method"],
                          "review_status": record["review_status"]})
    write_csv(output_dir / "koap_remarks.csv", rows, list(rows[0]))
    link_fields = ["remark_id", "source_sha256", "target_index", "category_id", "entity_id", "code_id",
                   "level", "article", "part", "paragraph", "chapter", "section", "remark_number",
                   "scope_qualifier", "annotation_method", "review_status"]
    write_csv(output_dir / "remark_targets.csv", links, link_fields)
    counts = Counter(r["category"] for r in records)
    write_categories(output_dir / "remark_categories.txt", counts)
    manifest = {"source_document": str(source), "source_sha256": source_hash,
                "edition_label": source.stem, "edition_source": "filename; not independently verified",
                "remarks": len(records), "remark_blocks": max(r["block_index"] for r in records),
                "active_text_rows": len(records) - counts[9], "repealed_placeholders": counts[9],
                "target_links": len(links), "categories": dict(sorted(counts.items())),
                "priority_review_rows": [r["remark_id"] for r in records if r["review_priority"] == "высокий"],
                "annotation_status": "provisional; expert review pending",
                "transitive_remark_links_expanded": False}
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fields = ("category", "targets", "annotation_method", "annotation_reason", "review_priority",
              "review_status", "candidate_targets", "text_sha256")
    annotation_file = {"source_sha256": source_hash,
                       "annotations": {r["remark_id"]: {k: r[k] for k in fields} for r in records}}
    (output_dir / "scope_annotations.json").write_text(json.dumps(annotation_file, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    workbook = Workbook()
    review = workbook.active
    review.title = "Примечания"
    headers = ["Примечание", "Категория", "Применяется к", "Расположение", "Основание разметки",
               "Приоритет проверки", "ID примечания", "Категория эксперта", "Сущности эксперта", "Комментарий эксперта"]
    review.append(headers)
    for row in rows:
        location = f"Статья {row['article']}" + (f", примечание {row['remark_number']}" if row["remark_number"] else ", без номера")
        review.append([row["remark_text"], row["category_label"], row["applies_to"], location,
                       row["annotation_reason"], row["review_priority"], row["remark_id"], "", "", ""])
    review.column_dimensions["A"].width = 105
    for column, width in {"B": 30, "C": 65, "D": 30, "E": 65, "F": 20, "G": 18,
                          "H": 22, "I": 60, "J": 60}.items():
        review.column_dimensions[column].width = width
    validation = DataValidation(type="whole", operator="between", formula1=0, formula2=9, allow_blank=True)
    validation.error = "Введите номер категории от 0 до 9."
    validation.showErrorMessage = True
    review.add_data_validation(validation)
    validation.add(f"H2:H{len(rows) + 1}")
    review["A1"].comment = Comment("Текст хранится целиком. Для длинных примечаний разверните строку формул или откройте ячейку.", "LCD")
    review["C1"].comment = Comment("Непосредственная область применения. Условия и исключения сохраняются в полном тексте. Все связи предварительные.", "LCD")
    colors = {0: "D9EAF7", 1: "E2F0D9", 2: "FFF2CC", 3: "FCE4D6", 4: "DDEBF7",
              5: "E4DFEC", 6: "E2D9F3", 7: "F4CCCC", 8: "F4B183", 9: "D9D9D9"}
    for index, record in enumerate(records, 2):
        review.cell(index, 2).fill = PatternFill("solid", fgColor=colors[record["category"]])
        if record["review_priority"] == "высокий":
            review.cell(index, 6).fill = PatternFill("solid", fgColor="F4B183")
        review.row_dimensions[index].height = min(409, max(60, 15 * (len(record["remark_text"]) // 100 + record["remark_text"].count("\n") + 2)))
        for col in (8, 9, 10):
            review.cell(index, col).fill = PatternFill("solid", fgColor="F2F2F2")
    targets = workbook.create_sheet("Связи")
    targets.append(link_fields)
    for link in links:
        targets.append([link[f] for f in link_fields])
    categories = workbook.create_sheet("Категории")
    categories.append(["Номер", "Категория", "Описание", "Число примечаний"])
    for number, (name, description) in CATEGORIES.items():
        categories.append([number, name, description, counts[number]])
    categories.column_dimensions["C"].width = 110
    metadata = workbook.create_sheet("Сведения")
    metadata.append(["Поле", "Значение"])
    for key, value in manifest.items():
        metadata.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else str(value)])
    metadata.append(["Проверка экспертом", "Все строки предварительные. Исправления: последние три колонки листа «Примечания». Сначала фильтруйте высокий приоритет."])
    metadata.append(["Связывание", "Лист «Связи» / remark_targets.csv: remark_id → entity_id. Используйте вместе с source_sha256. Кандидаты из category 8 не являются подтвержденными связями."])
    metadata.append(["Утратившие силу", "Категория 9: место размещения сохранено, действующие связи отсутствуют."])
    metadata.append(["Длинные примечания", "Полный текст сохранен в одной ячейке. Разверните строку формул, если высоты строки недостаточно."])
    metadata.column_dimensions["B"].width = 120
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.zoomScale = 80
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="203864")
            cell.font = Font(color="FFFFFF", bold=True)
        for line in sheet:
            for cell in line:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                # Legal texts/labels are always strings, including leading '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if sheet != review:
            for column in sheet.columns:
                letter = column[0].column_letter
                if sheet.column_dimensions[letter].width <= 13:
                    sheet.column_dimensions[letter].width = min(60, max(18, max(len(str(c.value or "")) for c in column) + 2))
    workbook.save(output_dir / "koap_remarks_review.xlsx")
    return manifest


def main() -> None:
    """Run extraction and annotation exports using explicit local paths."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--annotations", type=Path, help="Local curated JSON; source and text hashes must match")
    args = parser.parse_args()
    records = extract_remarks(read_paragraphs(args.source))
    overrides = {}
    if args.annotations:
        payload = json.loads(args.annotations.read_text(encoding="utf-8"))
        if payload["source_sha256"] != hashlib.sha256(args.source.read_bytes()).hexdigest():
            raise ValueError("Annotation source hash differs from the input DOCX")
        overrides = payload["annotations"]
    records = annotate_records(records, overrides)
    summary = export_records(records, args.source, args.output_dir)
    (args.output_dir / "extracted_remarks.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
