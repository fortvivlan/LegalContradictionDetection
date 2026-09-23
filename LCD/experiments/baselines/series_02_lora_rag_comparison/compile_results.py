"""Compile the three LoRA/RAG evaluation variants into one compact report."""

from __future__ import annotations

import argparse
import html
from collections import defaultdict
from pathlib import Path
from LCD.shared.common import REPOSITORY_ROOT, DEFAULT_RAG_DIR
from types import SimpleNamespace
from typing import Any, Mapping, Sequence

from LCD.shared.rag.evaluation import (
    _align_dataset,
    _annotations_by_hypothesis,
    _citation_key,
    _document_hypotheses,
    _merge_full_annotations,
    _retrieved_keys,
    read_rag_workbook,
)


DEFAULT_RESULT_PACKS = {
    "Baseline RAG": REPOSITORY_ROOT / "local/experiments/baselines/series_01_classifiers/full_pipeline_evaluation_baseline/results",
    "RAG-Qwen 100:60": Path(
        REPOSITORY_ROOT / "local/experiments/baselines/series_02_lora_rag_comparison/rag_qwen_legacy_lora/results"
    ),
    "RAG-Qwen 40:20": Path(
        REPOSITORY_ROOT / "local/experiments/baselines/series_02_lora_rag_comparison/rag_qwen_topk20/results"
    ),
}
LORA_NAMES = {
    "Qwen_Qwen3-8B": "Qwen3-8B",
    "meta-llama_Llama-3.1-8B": "Llama-3.1-8B",
    "mistralai_Ministral-8B-Instruct-2410": "Ministral-8B",
    "t-tech_T-lite-it-2.1": "T-lite-it-2.1",
}
LORA_ORDER = tuple(LORA_NAMES.values())
VARIANT_ORDER = tuple(DEFAULT_RESULT_PACKS)
SCORE_COLUMNS = (
    "accuracy",
    "macro_f1",
    "contradiction_precision",
    "contradiction_recall",
    "contradiction_f1",
)
COLORS = {
    "Baseline RAG": "#4C78A8",
    "RAG-Qwen 100:60": "#F58518",
    "RAG-Qwen 40:20": "#54A24B",
}


def _short_lora_name(model_name: str) -> str:
    prefix = "models__lora__"
    if not model_name.startswith(prefix):
        raise ValueError(f"Not a LoRA job name: {model_name}")
    remainder = model_name[len(prefix) :]
    for slug, short_name in LORA_NAMES.items():
        if remainder.startswith(f"{slug}__"):
            return short_name
    raise ValueError(f"Unknown LoRA job name: {model_name}")


def _read_lora_scores(result_packs: Mapping[str, Path]):
    import pandas as pd

    frames = []
    for variant, root in result_packs.items():
        score_paths = sorted(root.glob("jobs/models__lora__*/scores.csv"))
        if not score_paths:
            raise FileNotFoundError(f"No LoRA scores found in {root}")
        for score_path in score_paths:
            frame = pd.read_csv(score_path)
            frame = frame.loc[
                frame["evaluation_scope"].isin(("autotest_model", "autotest_total"))
                & frame["test_dataset"].eq("Full")
                & frame["task"].eq("ternary")
            ].copy()
            frame.insert(0, "rag_variant", variant)
            frames.append(frame)
    combined = pd.concat(frames, ignore_index=True)
    combined.insert(
        1, "lora_name", combined["model_name"].map(_short_lora_name)
    )
    combined["dataset"] = combined["test_dataset"].str.upper()
    compact = combined[
        [
            "rag_variant",
            "lora_name",
            "task",
            "evaluation_scope",
            "dataset",
            *SCORE_COLUMNS,
        ]
    ].copy()
    compact[list(SCORE_COLUMNS)] = compact[list(SCORE_COLUMNS)].round(2)
    compact["rag_variant"] = pd.Categorical(
        compact["rag_variant"], categories=VARIANT_ORDER, ordered=True
    )
    compact["lora_name"] = pd.Categorical(
        compact["lora_name"], categories=LORA_ORDER, ordered=True
    )
    compact.sort_values(
        ["task", "dataset", "evaluation_scope", "lora_name", "rag_variant"],
        inplace=True,
    )
    compact[["rag_variant", "lora_name"]] = compact[
        ["rag_variant", "lora_name"]
    ].astype(str)
    compact.reset_index(drop=True, inplace=True)
    return compact


def _aligned_rag_rows(rag_test_dir: Path, test_docx_dir: Path):
    full_path = rag_test_dir / "RAG_FULL_test.xlsx"
    additional_path = rag_test_dir / "RAG_FULL_additional_test.xlsx"
    full, _ = _align_dataset(
        "FULL",
        _document_hypotheses(test_docx_dir / "Full"),
        _merge_full_annotations(
            read_rag_workbook(full_path), read_rag_workbook(additional_path)
        ),
        (full_path, additional_path),
    )
    return {"FULL": full}


def _representative_results(result_pack: Path) -> Path:
    candidates = sorted(result_pack.glob("jobs/models__lora__*/results.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No LoRA result workbooks found in {result_pack}")
    return candidates[0]


def _recall_from_saved_pairs(result_pack: Path, aligned_rows):
    import pandas as pd

    workbook = _representative_results(result_pack)
    pairs = pd.read_excel(workbook, sheet_name="document_pairs")
    grouped = {
        (str(dataset).upper(), str(document), int(sentence_index)): group
        for (dataset, document, sentence_index), group in pairs.groupby(
            ["test_dataset", "document", "sentence_index"], dropna=False
        )
    }
    rows = []
    for dataset, occurrences in aligned_rows.items():
        counters: dict[str, list[int]] = defaultdict(lambda: [0, 0])
        for occurrence in occurrences:
            gold = {_citation_key(value) for value in occurrence["gold_citations"]}
            if not gold:
                continue
            group = grouped.get(
                (
                    dataset,
                    occurrence["document"],
                    int(occurrence["sentence_index"]),
                )
            )
            records: Sequence[Any] = ()
            methods: set[str] = set()
            if group is not None:
                records = tuple(
                    SimpleNamespace(source=str(row.source))
                    for row in group.itertuples(index=False)
                )
                methods = set(group["retrieval_method"].astype(str))
            system = "rules" if methods and methods <= {"exact"} else "faiss"
            hits = len(gold & _retrieved_keys(records))
            counters[system][0] += hits
            counters[system][1] += len(gold)
            counters["total"][0] += hits
            counters["total"][1] += len(gold)
        rows.append(
            {
                "dataset": dataset,
                **{
                    f"{system}_recall": (
                        counters[system][0] / counters[system][1]
                        if counters[system][1]
                        else None
                    )
                    for system in ("faiss", "rules", "total")
                },
            }
        )
    return pd.DataFrame(rows)


def _read_rag_recall(
    *,
    result_packs: Mapping[str, Path],
    sbert_recall_path: Path,
    embedding_recall_path: Path,
    rag_test_dir: Path,
    test_docx_dir: Path,
):
    import pandas as pd

    sbert = pd.read_csv(sbert_recall_path)
    baseline = sbert.loc[
        sbert["variant"].eq("baseline_embeddings__no_reranker")
        & sbert["candidate_top_k"].eq(40)
        & sbert["final_top_k"].eq(20)
    ]
    if len(baseline) != 1:
        raise ValueError("Expected one baseline RAG Recall row at 40:20")
    embedding = pd.read_csv(embedding_recall_path)
    qwen_100 = embedding.loc[
        embedding["model_alias"].eq("qwen3_embedding_0_6b")
        & embedding["candidate_top_k"].eq(100)
        & embedding["final_top_k"].eq(60)
    ]
    if len(qwen_100) != 1:
        raise ValueError("Expected one Qwen RAG Recall row at 100:60")

    aligned = _aligned_rag_rows(rag_test_dir, test_docx_dir)
    calculated_baseline = _recall_from_saved_pairs(
        result_packs["Baseline RAG"], aligned
    ).set_index("dataset")
    calculated_100 = _recall_from_saved_pairs(
        result_packs["RAG-Qwen 100:60"], aligned
    ).set_index("dataset")
    calculated_40 = _recall_from_saved_pairs(
        result_packs["RAG-Qwen 40:20"], aligned
    )
    known_100 = qwen_100.iloc[0]
    for dataset, prefix in (("FULL", "full"),):
        for system in ("faiss", "rules", "total"):
            calculated = calculated_baseline.loc[dataset, f"{system}_recall"]
            published = baseline.iloc[0][f"{prefix}_{system}_recall"]
            if abs(float(calculated) - float(published)) > 1e-12:
                raise ValueError(
                    "Saved-pair Recall did not reproduce published baseline "
                    f"{dataset} {system} Recall"
                )
            calculated = calculated_100.loc[dataset, f"{system}_recall"]
            published = known_100[f"{prefix}_{system}_recall"]
            if abs(float(calculated) - float(published)) > 1e-12:
                raise ValueError(
                    "Saved-pair Recall did not reproduce published Qwen 100:60 "
                    f"{dataset} {system} Recall"
                )

    recall_rows = []
    for variant, source, candidate, final in (
        ("Baseline RAG", baseline.iloc[0], 40, 20),
        ("RAG-Qwen 100:60", qwen_100.iloc[0], 100, 60),
    ):
        for dataset, prefix in (("FULL", "full"),):
            recall_rows.append(
                {
                    "rag_variant": variant,
                    "candidate_top_k": candidate,
                    "final_top_k": final,
                    "dataset": dataset,
                    "faiss_recall": source[f"{prefix}_faiss_recall"],
                    "rules_recall": source[f"{prefix}_rules_recall"],
                    "total_recall": source[f"{prefix}_total_recall"],
                }
            )
    for row in calculated_40.to_dict("records"):
        recall_rows.append(
            {
                "rag_variant": "RAG-Qwen 40:20",
                "candidate_top_k": 40,
                "final_top_k": 20,
                **row,
            }
        )
    recall = pd.DataFrame(recall_rows)
    recall[list(("faiss_recall", "rules_recall", "total_recall"))] = recall[
        ["faiss_recall", "rules_recall", "total_recall"]
    ].round(2)
    recall["rag_variant"] = pd.Categorical(
        recall["rag_variant"], categories=VARIANT_ORDER, ordered=True
    )
    recall.sort_values(["rag_variant", "dataset"], inplace=True)
    recall["rag_variant"] = recall["rag_variant"].astype(str)
    recall.reset_index(drop=True, inplace=True)
    return recall


def _svg_barplot(scores, task: str, output_path: Path) -> None:
    selected = scores.loc[
        scores["task"].eq(task)
        & scores["evaluation_scope"].eq("autotest_total")
    ]
    width, height = 1420, 500
    panel_width, panel_height = 650, 330
    panel_origins = ((80, 130), (760, 130))
    panels = (
        ("FULL", "contradiction_precision", "FULL — precision"),
        ("FULL", "contradiction_recall", "FULL — recall"),
    )
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="white"/>',
        f'<text x="710" y="42" text-anchor="middle" font-family="Arial" font-size="26" font-weight="bold">{html.escape(task.title())} — contradiction metrics (autotest_total)</text>',
    ]
    legend_x = 400
    for variant in VARIANT_ORDER:
        color = COLORS[variant]
        parts.append(
            f'<rect x="{legend_x}" y="67" width="18" height="18" fill="{color}"/>'
        )
        parts.append(
            f'<text x="{legend_x + 25}" y="82" font-family="Arial" font-size="15">{html.escape(variant)}</text>'
        )
        legend_x += 210

    for (dataset, metric, title), (left, top) in zip(panels, panel_origins):
        plot_left, plot_top = left + 55, top + 35
        plot_width, plot_height = panel_width - 75, panel_height - 85
        parts.append(
            f'<text x="{left + panel_width / 2}" y="{top + 18}" text-anchor="middle" font-family="Arial" font-size="18" font-weight="bold">{title}</text>'
        )
        for tick in range(0, 11, 2):
            value = tick / 10
            y = plot_top + plot_height * (1 - value)
            parts.append(
                f'<line x1="{plot_left}" y1="{y:.1f}" x2="{plot_left + plot_width}" y2="{y:.1f}" stroke="#dddddd" stroke-width="1"/>'
            )
            parts.append(
                f'<text x="{plot_left - 9}" y="{y + 5:.1f}" text-anchor="end" font-family="Arial" font-size="12">{value:.1f}</text>'
            )
        group_width = plot_width / len(LORA_ORDER)
        bar_width = group_width * 0.22
        for model_index, model in enumerate(LORA_ORDER):
            group_center = plot_left + group_width * (model_index + 0.5)
            model_rows = selected.loc[
                selected["dataset"].eq(dataset)
                & selected["lora_name"].eq(model)
            ]
            for variant_index, variant in enumerate(VARIANT_ORDER):
                values = model_rows.loc[model_rows["rag_variant"].eq(variant), metric]
                if values.empty:
                    continue
                value = float(values.iloc[0])
                x = group_center + (variant_index - 1) * bar_width - bar_width / 2
                bar_height = value * plot_height
                y = plot_top + plot_height - bar_height
                parts.append(
                    f'<rect x="{x:.1f}" y="{y:.1f}" width="{bar_width - 2:.1f}" height="{bar_height:.1f}" fill="{COLORS[variant]}"/>'
                )
                parts.append(
                    f'<text x="{x + (bar_width - 2) / 2:.1f}" y="{max(plot_top + 10, y - 5):.1f}" text-anchor="middle" font-family="Arial" font-size="11">{value:.2f}</text>'
                )
            parts.append(
                f'<text x="{group_center:.1f}" y="{plot_top + plot_height + 20:.1f}" text-anchor="middle" font-family="Arial" font-size="12">{html.escape(model)}</text>'
            )
        parts.append(
            f'<line x1="{plot_left}" y1="{plot_top + plot_height}" x2="{plot_left + plot_width}" y2="{plot_top + plot_height}" stroke="#333333"/>'
        )
    parts.append("</svg>")
    output_path.write_text("\n".join(parts), encoding="utf-8")


def _format_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(32, max(len(str(cell.value or "")) for cell in column) + 2)
            sheet.column_dimensions[letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
    workbook.save(path)


def compile_lora_rag_results(
    *,
    result_packs: Mapping[str, str | Path] = DEFAULT_RESULT_PACKS,
    sbert_recall_path: str | Path = REPOSITORY_ROOT / "local/experiments/lawbook/series_01_embedding_reranking/results/sbert_legal_v1/rag_recall.csv",
    embedding_recall_path: str | Path = REPOSITORY_ROOT / "local/experiments/lawbook/series_01_embedding_reranking/results/embedding_stage3/rag_recall.csv",
    rag_test_dir: str | Path = REPOSITORY_ROOT / "local/data/benchmarks/rag_tests",
    test_docx_dir: str | Path = REPOSITORY_ROOT / "local/data/benchmarks/test_docx",
    output_dir: str | Path = REPOSITORY_ROOT / "local/experiments/baselines/series_02_lora_rag_comparison/compilation/results",
) -> dict[str, Path]:
    """Write the compact three-class LoRA/RAG comparison workbook and plot."""
    import pandas as pd

    packs = {name: Path(path) for name, path in result_packs.items()}
    if tuple(packs) != VARIANT_ORDER:
        raise ValueError(f"result_packs must be ordered as {VARIANT_ORDER}")
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    scores = _read_lora_scores(packs)
    recall = _read_rag_recall(
        result_packs=packs,
        sbert_recall_path=Path(sbert_recall_path),
        embedding_recall_path=Path(embedding_recall_path),
        rag_test_dir=Path(rag_test_dir),
        test_docx_dir=Path(test_docx_dir),
    )
    workbook_path = output / "lora_rag_results.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        scores.to_excel(writer, sheet_name="lora_scores", index=False)
        recall.to_excel(writer, sheet_name="rag_recall", index=False)
    _format_workbook(workbook_path)
    ternary_plot = output / "contradiction_precision_recall.svg"
    _svg_barplot(scores, "ternary", ternary_plot)
    return {
        "workbook": workbook_path.resolve(),
        "ternary_plot": ternary_plot.resolve(),
    }


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=REPOSITORY_ROOT / "local/experiments/baselines/series_02_lora_rag_comparison/compilation/results",
    )
    return parser.parse_args()


if __name__ == "__main__":
    for name, path in compile_lora_rag_results(
        output_dir=_parse_args().output_dir
    ).items():
        print(f"{name}: {path}")
